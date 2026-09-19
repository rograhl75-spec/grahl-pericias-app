import pytest
from docx import Document

from grahl_app.parser import parse_pre_relatorio
from grahl_app.photos import build_photo_entry, decode_photo_base64, photo_exists, photo_identity_hash
from grahl_app.utils import generate_next_process_id, normalize_process_data, sanitize_output_filename, validate_process_data


def test_generate_next_process_id_ignores_malformed_ids():
    existing = ["Proc_01", "Proc_09", "Proc_x", "Outro", "Proc_0011a", "Proc_10"]
    assert generate_next_process_id(existing) == "Proc_11"


def test_normalize_process_data_merges_defaults_without_shared_lists():
    normalized = normalize_process_data({"reclamante_nome": None, "quadro_epis": None, "campo_fotos": None})
    assert normalized["reclamante_nome"] == ""
    assert normalized["quadro_epis"] == []
    assert normalized["campo_fotos"] == []

    normalized["quadro_epis"].append({"descricao": "Capacete", "ca": "1", "data_entrega": "", "obs": ""})
    fresh = normalize_process_data({})
    assert fresh["quadro_epis"] == []


def test_validate_process_data_allows_partial_save_but_flags_generation_requirements():
    permissive = validate_process_data({"reclamante_nome": "Fulano"}, for_generation=False)
    result = validate_process_data({"reclamante_nome": "Fulano"}, for_generation=True)

    assert permissive.is_valid
    assert not result.is_valid
    assert set(result.missing_required_fields) == {"reclamada_nome", "processo_num"}


def test_parse_pre_relatorio_extracts_multiline_fields_and_epis():
    doc = Document()
    doc.add_paragraph("Número do Processo: 0001234-55.2024.5.09.0001")
    doc.add_paragraph("Reclamante (Autor/Autora): João da Silva (CPF: 123.456.789-00)")
    doc.add_paragraph("Reclamada (Ré / Empresa): Empresa XPTO (CNPJ: 11.222.333/0001-44)")
    doc.add_paragraph("Atividades Descritas: Primeira linha")
    doc.add_paragraph("Segunda linha da atividade")
    table = doc.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "Descrição do EPI"
    table.rows[0].cells[1].text = "C.A."
    table.rows[0].cells[2].text = "Data"
    table.rows[0].cells[3].text = "Obs"
    table.rows[1].cells[0].text = "Luva nitrílica"
    table.rows[1].cells[1].text = "12345"
    table.rows[1].cells[2].text = "Mensal"
    table.rows[1].cells[3].text = "Em bom estado"

    parsed = parse_pre_relatorio(doc)

    assert parsed["processo_num"] == "0001234-55.2024.5.09.0001"
    assert parsed["reclamante_nome"] == "João da Silva"
    assert parsed["reclamante_cpf"] == "123.456.789-00"
    assert parsed["reclamada_nome"] == "Empresa XPTO"
    assert parsed["reclamada_cnpj"] == "11.222.333/0001-44"
    assert parsed["atividades_inicial"] == "Primeira linha\nSegunda linha da atividade"
    assert parsed["quadro_epis"][0]["descricao"] == "Luva nitrílica"


def test_photo_hash_helpers_detect_duplicates_and_invalid_base64():
    entry = build_photo_entry(b"fake-image-bytes", gps="1,2")
    duplicate = build_photo_entry(b"fake-image-bytes", gps="3,4")
    assert photo_identity_hash(b"fake-image-bytes") == entry["hash"]
    assert photo_exists([entry], duplicate)
    assert decode_photo_base64(entry["base64"]) == b"fake-image-bytes"
    with pytest.raises(ValueError):
        decode_photo_base64("not-base64!!")


def test_sanitize_output_filename_removes_invalid_characters():
    filename = sanitize_output_filename('Ana / Maria:*? <teste>\nfinal')
    assert filename == "Documento_Oficial_Ana_Maria_teste_final.docx"
