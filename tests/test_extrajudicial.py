import ast
import base64
import copy
import csv
import io
import json
import re
from pathlib import Path
from unittest import mock

import pytest
import streamlit as st
from docx import Document
from docx.document import Document as DocxDocument
from PIL import Image
from streamlit.testing.v1 import AppTest

from core import database
from core.config import criar_dados_padrao
from core.extrajudicial import (
    CNIS_COLUMNS, DOCX_MARKER, ENV_COLUMNS, UNKNOWN, apply_import, export_payload,
    final_errors, invalidate_review, legacy_state, new_defaults, parse_cnis_csv,
    parse_docx, parse_json, recognized_import, review, validate,
)
from core.extrajudicial_document import build_extrajudicial_bytes
from grahl_app.documents import build_document_bytes

APP_FILE = Path(__file__).resolve().parents[1] / "app.py"
MODULE = "Laudo Extrajudicial Previdenciário (LTCAT+PPP)"


@pytest.fixture
def active_parser():
    tree = ast.parse(APP_FILE.read_text())
    node = next(node for node in tree.body
                if isinstance(node, ast.FunctionDef) and node.name == "parse_pre_relatorio")
    namespace = {"re": re}
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(APP_FILE), "exec"), namespace)
    return namespace["parse_pre_relatorio"]


@pytest.fixture
def complete_case():
    data = criar_dados_padrao()
    data.update(new_defaults())
    data.update({
        "modulo_atuacao": MODULE, "_revision": 7, "custom_legacy": "Preservar",
        "reclamante_nome": "Segurado Exemplo", "reclamante_cpf": "123",
        "reclamada_nome": "Oficina", "natureza_laudo": "Individual",
        "modo_temporal": "Contemporâneo", "responsavel_nome": "Profissional Exemplo",
        "responsavel_conselho": "CREA / PR / 123", "responsavel_art": "ART 456",
        "doc_ltcat": "Método profissional documentado",
        "conclusao_tecnica": "Não há elementos suficientes para caracterizar tempo especial.",
        "fontes_documentais": "CNIS e documentos conferidos em 01/10/2026.",
        "agentes_observados": "Ruído da máquina, observação direta.",
        "profissiografia_real": "Opera máquinas e realiza manutenção.",
        "campo_data": "01/10/2026", "campo_horario": "09:30",
        "local_diligencia": "Oficina de exemplo", "presentes_pericia": "Segurado e acompanhante",
        "campo_declaracoes_autor": "Declaração exclusiva do segurado",
        "campo_declaracoes_reu": "Declaração exclusiva do empregador",
        "campo_medicoes": "Dosimetria registrada, certificado em anexo",
        "cnis_periodos": [{
            "periodo_id": "P1", "inicio": "2020-01-01", "fim": "2021-12-31",
            "empresa": "Oficina", "cnpj": "11.222.333/0001-44", "condicao": "CLT",
            "cargo": "Mecânico", "setor": "Manutenção", "atividade": "Reparo de máquinas",
        }],
        "avaliacoes_ambientais": [{
            **dict.fromkeys(ENV_COLUMNS, ""),
            "periodo_id": "P1", "tipo": "F", "agente": "Ruído", "intensidade": "82",
            "unidade": "dB(A)", "metodo": "NHO-01", "fonte": "Máquina",
            "via": "Auditiva", "exposicao": "Intermitente",
            "epc": UNKNOWN, "epi": "Não", "ca": "12345",
            "responsavel": "Profissional Exemplo", "conselho": "CREA / PR / 123",
            "instrumento": "Dosímetro XYZ", "calibracao": "2026-09-01",
        }],
    })
    return data


def document_text(doc):
    return "\n".join([p.text for p in doc.paragraphs] + [
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ])


def test_new_defaults_are_unknown_and_independent():
    data = new_defaults()
    assert data["apr_fisicos"] == data["extemp_justificativa"] == data["conclusao_tecnica"] == ""
    assert data["extemp_layout"] is None
    assert data["extemp_layout_estado"] == UNKNOWN
    assert not data["revisado_profissional"]
    data["cnis_periodos"].append({})
    assert new_defaults()["cnis_periodos"] == []
    assert validate(new_defaults()) == []


def test_legacy_persistence_preserves_existing_content_and_judicial_defaults():
    original = {"modulo_atuacao": MODULE, "extemp_layout": False,
                "extemp_maquinas": True, "extemp_justificativa": "Texto legado",
                "data_admissao": "1999 até 2003", "apr_fisicos": "APR original"}
    normalized = database._normalizar_dados_processo(original)
    assert original == {k: normalized[k] for k in original}
    assert legacy_state(normalized, "extemp_layout") == "Não"
    assert legacy_state(normalized, "extemp_maquinas") == "Sim"
    assert legacy_state(normalized, "extemp_epc") == UNKNOWN
    assert normalized["apr_quimicos"] == ""
    judicial = database._normalizar_dados_processo({})
    assert judicial["apr_fisicos"] == criar_dados_padrao()["apr_fisicos"]


@pytest.mark.parametrize("field,value,expected", [
    ("inicio", "31/01/2020", "datas"),
    ("inicio", "2020-02-30", "datas"),
    ("fim", "2019-12-31", "anterior"),
    ("empresa", "", "empresa"),
    ("cnpj", "123", "CNPJ"),
    ("condicao", "talvez", "condição"),
    ("periodo_id", "", "periodo_id"),
])
def test_period_validation(complete_case, field, value, expected):
    complete_case["cnis_periodos"][0][field] = value
    assert any(expected in error for error in validate(complete_case))


def test_period_id_unique_and_open_ended_allowed(complete_case):
    complete_case["cnis_periodos"][0]["fim"] = ""
    assert validate(complete_case) == []
    complete_case["cnis_periodos"].append(copy.deepcopy(complete_case["cnis_periodos"][0]))
    assert any("único" in error for error in validate(complete_case))


@pytest.mark.parametrize("field,value", [("periodo_id", "P99"), ("tipo", "X"), ("epi", ""),
                                        ("fonte", ""), ("metodo", "")])
def test_environmental_validation(complete_case, field, value):
    complete_case["avaliacoes_ambientais"][0][field] = value
    assert validate(complete_case)


def test_final_gate_requires_conclusion_and_explicit_review(complete_case):
    assert final_errors(complete_case)
    assert not final_errors(review(complete_case))
    complete_case["conclusao_tecnica"] = ""
    with pytest.raises(ValueError, match="Conclusão"):
        review(complete_case)


@pytest.mark.parametrize("field,value", [
    ("campo_medicoes", "Alteradas"), ("campo_fotos", [{"base64": "AA==", "legenda": "Nova foto"}]),
    ("conclusao_tecnica", "Conclusão alterada"), ("extemp_evidencias", "Outra fonte"),
])
def test_substantive_changes_invalidate_review(complete_case, field, value):
    reviewed = review(complete_case)
    reviewed[field] = value
    assert final_errors(reviewed)
    updated = invalidate_review(reviewed)
    assert updated["revisado_profissional"] is False
    assert updated["_revision"] == 7


def test_revision_changes_do_not_invalidate_evidence_review(complete_case):
    reviewed = review(complete_case)
    reviewed["_revision"] += 1
    assert not final_errors(reviewed)


def test_extemporaneous_requires_historical_evidence(complete_case):
    complete_case["modo_temporal"] = "Extemporâneo"
    assert any("Evidências" in error for error in validate(complete_case, final=True))
    complete_case.update(extemp_evidencias="Documentos históricos comparados por período",
                         extemp_justificativa="Limitações da comparação documentadas")
    assert not final_errors(review(complete_case))


def test_csv_exact_schema_and_validation(complete_case):
    stream = io.StringIO()
    writer = csv.DictWriter(stream, fieldnames=CNIS_COLUMNS)
    writer.writeheader()
    writer.writerows(complete_case["cnis_periodos"])
    assert parse_cnis_csv(stream.getvalue().encode("utf-8-sig")) == complete_case["cnis_periodos"]
    for bad in ("inicio,fim\n2020-01-01,2021-01-01", ",".join(CNIS_COLUMNS),
                stream.getvalue().replace("2020-01-01", "2020-99-99"),
                stream.getvalue() + "extra,columns\n"):
        with pytest.raises(ValueError):
            parse_cnis_csv(bad)


def test_reviewed_import_preserves_metadata_photos_and_false_empty_values(complete_case):
    current = review(complete_case)
    current["campo_fotos"] = [{"base64": "preserve", "legenda": "Preserve"}]
    incoming = {"conclusao_tecnica": "", "extemp_layout": False, "reclamante_nome": "Outro",
                "_revision": 99, "id": "other", "campo_fotos": []}
    snapshot = copy.deepcopy(current)
    with pytest.raises(ValueError, match="Confirme"):
        apply_import(current, incoming, ["conclusao_tecnica"])
    updated = apply_import(current, incoming, ["conclusao_tecnica", "extemp_layout"], confirmed=True)
    assert updated["conclusao_tecnica"] == ""
    assert updated["extemp_layout"] is False
    assert updated["_revision"] == 7
    assert updated["campo_fotos"] == current["campo_fotos"]
    assert updated["reclamante_nome"] == current["reclamante_nome"]
    assert updated["custom_legacy"] == "Preservar"
    assert current == snapshot
    assert not updated["revisado_profissional"]
    with pytest.raises(ValueError):
        apply_import(current, incoming, ["id"], confirmed=True)


@pytest.mark.parametrize("incoming", [
    {"cnis_periodos": "bad"}, {"campo_medicoes": False}, {"extemp_layout": "false"},
    {"reclamante_nome": {"nested": "bad"}}, {"campo_fotos": []},
    {"cnis_periodos": [{"periodo_id": "P1", "inicio": []}]},
])
def test_malformed_imports_rejected(incoming):
    with pytest.raises(ValueError):
        recognized_import(incoming)


def test_document_and_json_roundtrip_with_active_parser(complete_case, active_parser):
    reviewed = review(complete_case)
    json_data = parse_json(json.dumps(export_payload(reviewed)))
    buffer, warnings = build_extrajudicial_bytes(reviewed)
    assert not warnings
    imported = parse_docx(Document(buffer), active_parser)
    assert imported == json_data
    assert imported["cnis_periodos"] == complete_case["cnis_periodos"]
    assert "_revision" not in imported
    assert "revisado_profissional" not in imported
    assert "campo_fotos" not in imported


def test_active_legacy_parser_import_is_recognized_and_not_applied(active_parser):
    doc = Document()
    doc.add_paragraph("[START_RECLAMANTE_NOME]Nome importado[END_RECLAMANTE_NOME]")
    doc.add_paragraph("[START_APR_FISICOS]Ruído a verificar[END_APR_FISICOS]")
    incoming = parse_docx(doc, active_parser)
    assert incoming["reclamante_nome"] == "Nome importado"
    assert incoming["apr_fisicos"] == "Ruído a verificar"
    assert "modulo_atuacao" not in incoming
    doc = Document()
    doc.add_paragraph(DOCX_MARKER)
    doc.add_paragraph("{bad json")
    with pytest.raises(ValueError, match="JSON"):
        parse_docx(doc, active_parser)


def test_report_includes_inspection_ppp_and_professional_conclusion(complete_case):
    image = io.BytesIO()
    Image.new("RGB", (8, 8), "blue").save(image, format="PNG")
    complete_case["campo_fotos"] = [{"base64": base64.b64encode(image.getvalue()).decode(),
                                     "legenda": "Máquina inspecionada", "gps": "-23,-51"}]
    reviewed = review(complete_case)
    for builder in (build_extrajudicial_bytes, build_document_bytes):
        buffer, warnings = builder(reviewed)
        assert warnings == []
        doc = Document(buffer)
        content = document_text(doc)
        for expected in ("LTCAT", "CNIS", "P1", "Reparo de máquinas", "NHO-01",
                         "82", "dB(A)", "Dosímetro XYZ", "12345", "ART 456",
                         "Declaração exclusiva do segurado", "Declaração exclusiva do empregador",
                         "Dosimetria registrada", "Oficina de exemplo", "09:30",
                         "NÃO É PPP OFICIAL ASSINADO", "Máquina inspecionada", "-23,-51"):
            assert expected in content
        assert complete_case["conclusao_tecnica"] in content
        assert "Conclui-se que as atividades" not in content
        assert len(doc.inline_shapes) >= 1


def test_draft_never_claims_review_or_positive_conclusion():
    buffer, warnings = build_extrajudicial_bytes(new_defaults(), preliminary=True)
    content = document_text(Document(buffer))
    assert "RASCUNHO" in content
    assert "CHECKLIST" in content
    assert "Mudança de layout: Não informado" in content
    assert "preenchendo os requisitos" not in content
    assert "mantêm-se inalteradas" not in content
    assert not warnings
    with pytest.raises(ValueError):
        build_extrajudicial_bytes(new_defaults())


def test_judicial_document_does_not_gain_extrajudicial_sections():
    data = criar_dados_padrao()
    data.update(reclamante_nome="Autor judicial", processo_num="CNJ judicial",
                quesitos_juizo="Quesito do juízo", campo_medicoes="Medição judicial")
    buffer, warnings = build_document_bytes(data)
    content = document_text(Document(buffer))
    assert "CNJ judicial" in content
    assert "Quesito do juízo" in content
    assert "Medição judicial" in content
    assert "DADOS ESTRUTURADOS PARA REIMPORTAÇÃO" not in content
    assert not warnings


@pytest.fixture
def app_case(complete_case):
    records = {"Proc_01": complete_case}
    with mock.patch.object(database, "carregar_dados", side_effect=lambda: copy.deepcopy(records)):
        app = AppTest.from_file(str(APP_FILE)).run()
        yield app, records


def test_extrajudicial_app_has_stable_controls_and_partial_save(app_case):
    app, records = app_case
    app.sidebar.button[1].click().run()
    assert not app.exception
    assert len(app.tabs) == 4
    assert len(app.get("dataframe")) == 3
    assert len(app.text_input) + len(app.text_area) > 25
    assert app.text_area(key="extra_Proc_01_conclusao_tecnica").value == records["Proc_01"]["conclusao_tecnica"]
    original = copy.deepcopy(records["Proc_01"])
    app.text_area(key="extra_Proc_01_conclusao_tecnica").set_value("")
    assert records["Proc_01"] == original

    def save(ident, data):
        records[ident] = copy.deepcopy(data)
        return True

    with mock.patch.object(database, "salvar_processo", side_effect=save) as persist:
        next(button for button in app.button if "Salvar rascunho" in button.label).click().run()
    assert not app.exception
    persist.assert_called_once()
    assert records["Proc_01"]["conclusao_tecnica"] == ""
    assert records["Proc_01"]["_revision"] == 7
    assert not records["Proc_01"]["revisado_profissional"]
    assert records["Proc_01"]["custom_legacy"] == "Preservar"


def test_app_review_and_final_export_gate(app_case):
    app, records = app_case
    app.sidebar.button[4].click().run()
    assert not app.exception
    assert app.button(key="gerar_documento_Proc_01").disabled
    app.sidebar.button[1].click().run()
    assert app.button(key="extra_Proc_01_revisar").disabled

    def save(ident, data):
        records[ident] = copy.deepcopy(data)
        return True

    app.checkbox(key="extra_Proc_01_confirmar_revisao").check().run()
    with mock.patch.object(database, "salvar_processo", side_effect=save):
        app.button(key="extra_Proc_01_revisar").click().run()
    assert not app.exception
    assert not final_errors(records["Proc_01"])
    app.sidebar.button[4].click().run()
    assert not app.button(key="gerar_documento_Proc_01").disabled
    documents = []
    original_save = DocxDocument.save

    def capture(doc, stream):
        documents.append(doc)
        original_save(doc, stream)

    with mock.patch.object(DocxDocument, "save", autospec=True, side_effect=capture):
        app.button(key="gerar_documento_Proc_01").click().run()
    assert not app.exception
    assert len(documents) == 1
    content = document_text(documents[0])
    assert "Declaração exclusiva do segurado" in content
    assert "NÃO É PPP OFICIAL ASSINADO" in content
    assert records["Proc_01"]["conclusao_tecnica"] in content
    records["Proc_01"]["campo_medicoes"] = "Evidência alterada"
    app.run()
    assert app.button(key="gerar_documento_Proc_01").disabled


def test_extrajudicial_shared_field_screen_and_draft_export(app_case):
    app, records = app_case
    app.sidebar.button[2].click().run()
    assert not app.exception
    assert app.text_input(key="cd_Proc_01").value == "01/10/2026"
    app.sidebar.button[4].click().run()
    app.button(key="extra_Proc_01_preliminar").click().run()
    assert not app.exception
    assert len(app.get("download_button")) == 1
    assert not records["Proc_01"]["revisado_profissional"]


def test_shared_field_save_invalidates_review_without_losing_structured_data(app_case):
    app, records = app_case
    records["Proc_01"] = review(records["Proc_01"])
    original_periods = copy.deepcopy(records["Proc_01"]["cnis_periodos"])
    app.sidebar.button[2].click().run()
    app.text_area(key="cm_Proc_01").set_value("Nova medição documentada")

    def save(ident, data):
        records[ident] = copy.deepcopy(data)
        return True

    with mock.patch.object(database, "salvar_processo", side_effect=save) as persist:
        next(button for button in app.button if "Salvar Textos" in button.label).click().run()
    assert not app.exception
    persist.assert_called_once()
    assert records["Proc_01"]["campo_medicoes"] == "Nova medição documentada"
    assert records["Proc_01"]["cnis_periodos"] == original_periods
    assert records["Proc_01"]["_revision"] == 7
    assert not records["Proc_01"]["revisado_profissional"]
    app.sidebar.button[4].click().run()
    assert app.button(key="gerar_documento_Proc_01").disabled


def test_app_import_requires_selected_fields_and_confirmation(app_case):
    app, records = app_case
    upload = io.BytesIO(json.dumps(export_payload({"reclamante_nome": "Importado",
                                                 "conclusao_tecnica": ""})).encode())
    upload.name = "dados.json"
    original_uploader = st.file_uploader
    original = copy.deepcopy(records["Proc_01"])

    def uploader(*args, **kwargs):
        if kwargs.get("key") == "extra_Proc_01_import":
            return upload
        return original_uploader(*args, **kwargs)

    def save(ident, data):
        records[ident] = copy.deepcopy(data)
        return True

    with mock.patch("streamlit.file_uploader", side_effect=uploader):
        app.sidebar.button[1].click().run()
        assert not app.exception
        assert records["Proc_01"] == original
        apply_button = next(button for button in app.button if button.label == "Aplicar importação revisada")
        assert apply_button.disabled
        selected = next(select for select in app.multiselect if select.label == "Campos a aplicar")
        selected.set_value(["reclamante_nome"]).run()
        confirm = next(item for item in app.checkbox if "substituição" in item.label)
        confirm.check().run()
        with mock.patch.object(database, "salvar_processo", side_effect=save) as persist:
            next(button for button in app.button if button.label == "Aplicar importação revisada").click().run()
        assert not app.exception
    persist.assert_called_once()
    assert records["Proc_01"]["reclamante_nome"] == "Importado"
    assert records["Proc_01"]["conclusao_tecnica"] == original["conclusao_tecnica"]
    assert records["Proc_01"]["_revision"] == 7
    assert records["Proc_01"]["custom_legacy"] == "Preservar"


def test_app_creation_docx_preview_does_not_overwrite_typed_values(active_parser):
    doc = Document()
    doc.add_paragraph("[START_RECLAMANTE_NOME]Nome importado[END_RECLAMANTE_NOME]")
    upload = io.BytesIO()
    doc.save(upload)
    upload.name = "dados.docx"
    upload.seek(0)
    original_uploader = st.file_uploader

    def uploader(*args, **kwargs):
        if kwargs.get("key") == "uploader_0":
            return upload
        return original_uploader(*args, **kwargs)

    with mock.patch.object(database, "carregar_dados", return_value={}):
        app = AppTest.from_file(str(APP_FILE)).run()
        app.radio[0].set_value(app.radio[0].options[1]).run()
        app.text_input(key="novo_nome_0").set_value("Nome digitado")
        with mock.patch.object(database, "gerar_proximo_id", return_value="Proc_01"), \
             mock.patch.object(database, "salvar_processo", return_value=False):
            app.main.button[0].click().run()
        assert app.text_input(key="novo_nome_0").value == "Nome digitado"
        with mock.patch("streamlit.file_uploader", side_effect=uploader), \
             mock.patch.object(database, "salvar_processo") as persist:
            app.run()
            assert not app.exception
            assert app.text_input(key="novo_nome_0").value == "Nome digitado"
            assert app.session_state["parsed_data"] == {}
            next(button for button in app.button if "Usar campos" in button.label).click().run()
            assert not app.exception
            assert app.text_input(key="novo_nome_0").value == "Nome importado"
            assert app.session_state["parsed_data"]["reclamante_nome"] == "Nome importado"
            persist.assert_not_called()
