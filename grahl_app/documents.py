from __future__ import annotations

import io
import os
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor

from .constants import LOGO_FILE
from .photos import decode_photo_base64
from .utils import normalize_epi_list, normalize_process_data, safe_str

PRIMARY_COLOR = RGBColor(27, 54, 93)
FONT_NAME = "Abadi"


def _apply_run_style(run, *, size: float = 11, bold: bool = False, italic: bool = False, color=None) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _new_paragraph(doc: Document, *, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing: float = 1.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = spacing
    return paragraph


def _add_blank_line(doc: Document) -> None:
    _new_paragraph(doc)


def _add_bullet_inline(doc: Document, title: str, text: object) -> None:
    paragraph = _new_paragraph(doc)
    _apply_run_style(paragraph.add_run(f"• {safe_str(title)}: "), bold=True)
    _apply_run_style(paragraph.add_run(safe_str(text) or "[Informação não localizada]"))


def _add_topic_block(doc: Document, title: str, text: object) -> None:
    title_paragraph = _new_paragraph(doc)
    _apply_run_style(title_paragraph.add_run(f"• {safe_str(title)}:"), bold=True)
    content = safe_str(text) or "[Informação/Documento não localizado nos autos anexados]"
    for line in content.split("\n"):
        if line.strip():
            paragraph = _new_paragraph(doc)
            _apply_run_style(paragraph.add_run(line.strip()))


def _add_title(doc: Document, text: str, *, level: int = 2) -> None:
    heading = doc.add_heading(level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    heading.paragraph_format.space_before = Pt(18)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.line_spacing = 1.5
    run = heading.add_run(text)
    _apply_run_style(run, size=14 if level <= 2 else 12, bold=True, color=PRIMARY_COLOR)


def _set_table_width_100(table) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "5000")
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)


def _style_table(table, *, body_font_size: float = 9, striped_fill: str = "F2F2F2") -> None:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    _apply_run_style(run, size=body_font_size)
            if row_index == 0:
                _set_cell_shading(tc_pr, "1B365D")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _apply_run_style(run, size=body_font_size, bold=True, color=RGBColor(255, 255, 255))
            elif row_index % 2 == 0:
                _set_cell_shading(tc_pr, striped_fill)


def _set_cell_shading(tc_pr, fill: str) -> None:
    for child in list(tc_pr):
        if child.tag == qn("w:shd"):
            tc_pr.remove(child)
    tc_pr.append(parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), fill)))


def _add_key_value_table(doc: Document, rows: Iterable[tuple[str, object]]) -> None:
    rows = list(rows)
    table = doc.add_table(rows=len(rows) + 1, cols=2)
    _set_table_width_100(table)
    table.rows[0].cells[0].text = "Parâmetro"
    table.rows[0].cells[1].text = "Descrição"
    for index, (label, value) in enumerate(rows, start=1):
        table.rows[index].cells[0].text = safe_str(label)
        table.rows[index].cells[1].text = safe_str(value) or "[Informação não localizada]"
    _style_table(table)


def _add_epi_table(doc: Document, entries: list[dict[str, str]], *, periodicity_label: str, empty_message: str, observations_label: str) -> None:
    normalized_entries = normalize_epi_list(entries)
    if not normalized_entries:
        paragraph = _new_paragraph(doc)
        _apply_run_style(paragraph.add_run(empty_message))
        return

    table = doc.add_table(rows=len(normalized_entries) + 1, cols=4)
    _set_table_width_100(table)
    headers = ["Descrição do EPI", "C.A.", periodicity_label, observations_label]
    for header_index, header in enumerate(headers):
        table.rows[0].cells[header_index].text = header

    for row_index, epi in enumerate(normalized_entries, start=1):
        table.rows[row_index].cells[0].text = safe_str(epi.get("descricao"))
        table.rows[row_index].cells[1].text = safe_str(epi.get("ca"))
        table.rows[row_index].cells[2].text = safe_str(epi.get("data_entrega"))
        table.rows[row_index].cells[3].text = safe_str(epi.get("obs"))
    _style_table(table, body_font_size=8.5)


def _add_logo(doc: Document) -> None:
    if not os.path.exists(LOGO_FILE):
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(LOGO_FILE, width=Inches(2.0))
    _add_blank_line(doc)


def _build_judicial_document(doc: Document, process: dict) -> None:
    _add_title(doc, "1. IDENTIFICAÇÃO DO PROCESSO", level=2)
    _add_key_value_table(
        doc,
        [
            ("Número do Processo", process.get("processo_num", "")),
            ("Órgão Julgador / Vara", process.get("orgao_julgador", "")),
            ("Data de Autuação (Ajuizamento)", process.get("data_autuacao", "")),
            ("Valor da Causa", process.get("valor_causa", "")),
            ("Rito Processual", process.get("rito_processual", "")),
        ],
    )
    _add_blank_line(doc)

    _add_title(doc, "2. QUALIFICAÇÃO DAS PARTES", level=2)
    _add_bullet_inline(doc, "Reclamante (Autor/Autora)", f"{safe_str(process.get('reclamante_nome'))} (CPF: {safe_str(process.get('reclamante_cpf'))})")
    _add_bullet_inline(doc, "Advogados", process.get("reclamante_adv", ""))
    _add_bullet_inline(doc, "Reclamada (Ré/Empresa)", f"{safe_str(process.get('reclamada_nome'))} (CNPJ: {safe_str(process.get('reclamada_cnpj'))})")
    _add_bullet_inline(doc, "Advogados", process.get("reclamada_adv", ""))

    _add_title(doc, "3. DADOS DO CONTRATO DE TRABALHO", level=2)
    _add_bullet_inline(doc, "Data de Admissão", process.get("data_admissao", ""))
    _add_bullet_inline(doc, "Status do Contrato", process.get("status_contrato", ""))
    _add_bullet_inline(doc, "Período Imprescrito", process.get("periodo_imprescrito", ""))
    _add_bullet_inline(doc, "Cargo(s) / Função(ões)", process.get("cargos", ""))
    _add_bullet_inline(doc, "Setor / Lotação / Local", process.get("setor", ""))
    _add_bullet_inline(doc, "Última Remuneração", process.get("ultima_remuneracao", ""))

    _add_title(doc, "4. SÍNTESE TÉCNICA DA PETIÇÃO INICIAL", level=2)
    _add_bullet_inline(doc, "Objeto da Perícia", process.get("objeto_pericia", ""))
    _add_topic_block(doc, "Atividades Descritas", process.get("atividades_inicial", ""))
    _add_topic_block(doc, "Agentes Nocivos / Riscos Alegados", process.get("agentes_alegados", ""))
    _add_topic_block(doc, "Pedidos Técnicos", process.get("pedidos_tecnicos", ""))

    _add_title(doc, "5. SÍNTESE TÉCNICA DA CONTESTAÇÃO", level=2)
    _add_topic_block(doc, "Preliminares Periciais", process.get("preliminares_periciais", ""))
    _add_topic_block(doc, "Defesa de Mérito (SST)", process.get("defesa_merito_sst", ""))

    _add_title(doc, "6. STATUS E DADOS DA VISTORIA", level=2)
    _add_bullet_inline(doc, "Fase Processual Atual", process.get("fase_processual", ""))
    _add_bullet_inline(doc, "Data da Vistoria", process.get("campo_data", ""))
    _add_bullet_inline(doc, "Horário", process.get("campo_horario", ""))
    _add_bullet_inline(doc, "Local / Endereço", process.get("local_diligencia", ""))

    _add_title(doc, "7. ANÁLISE DOS DOCUMENTOS DE SST NOS AUTOS", level=2)
    _add_topic_block(doc, "LTCAT", process.get("doc_ltcat", ""))
    _add_topic_block(doc, "Laudo de Insalubridade / Periculosidade", process.get("doc_laudo", ""))
    _add_topic_block(doc, "PPP (Perfil Profissiográfico Previdenciário)", process.get("doc_ppp", ""))
    _add_topic_block(doc, "PGR / PPRA / PCMAT", process.get("doc_pgr", ""))
    _add_topic_block(doc, "Ordens de Serviço (OS) / Treinamentos", process.get("doc_os", ""))
    _add_topic_block(doc, "ASOs / PCMSO", process.get("doc_asos", ""))
    _add_topic_block(doc, "Outros Documentos Relevantes", process.get("doc_outros", ""))

    _add_title(doc, "8. FORNECIMENTO DE EQUIPAMENTOS DE PROTEÇÃO (EPIs)", level=2)
    _add_epi_table(
        doc,
        process.get("quadro_epis", []),
        periodicity_label="Data de Entrega",
        empty_message="[Nenhum EPI cadastrado para este processo]",
        observations_label="Observações",
    )
    _add_blank_line(doc)
    _add_topic_block(doc, "Síntese e Análise Crítica de EPIs", process.get("analise_epis_critica", ""))

    _add_title(doc, "9. QUESITOS FORMULADOS PARA A PERÍCIA", level=2)
    _add_topic_block(doc, "9.1. Quesitos do Juízo", process.get("quesitos_juizo", ""))
    _add_topic_block(doc, "9.2. Quesitos do Reclamante", process.get("quesitos_autor", ""))
    _add_topic_block(doc, "9.3. Quesitos da Reclamada", process.get("quesitos_reu", ""))

    _add_title(doc, "10. LEVANTAMENTOS DE CAMPO (DILIGÊNCIA & EVIDÊNCIAS)", level=2)
    _add_topic_block(doc, "Pessoas Presentes na Vistoria", process.get("presentes_pericia", ""))
    _add_topic_block(doc, "Informações prestadas pelo Autor", process.get("campo_declaracoes_autor", ""))
    _add_topic_block(doc, "Informações prestadas pelo Ré", process.get("campo_declaracoes_reu", ""))
    _add_topic_block(doc, "Medições Realizadas", process.get("campo_medicoes", ""))


def _build_previdenciario_document(doc: Document, process: dict) -> None:
    from core.extrajudicial_document import append_extrajudicial

    append_extrajudicial(doc, process)


def build_document_bytes(process_data: dict) -> tuple[io.BytesIO, list[str]]:
    if "Previdenciário" in safe_str(process_data.get("modulo_atuacao")):
        from core.extrajudicial_document import build_extrajudicial_bytes

        return build_extrajudicial_bytes(process_data)
    process = normalize_process_data(process_data)
    warnings: list[str] = []
    document = Document()
    for section in document.sections:
        section.top_margin = Cm(3.0)
        section.left_margin = Cm(3.0)
        section.bottom_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    _add_logo(document)
    title_paragraph = _new_paragraph(document, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    title = "LAUDO TÉCNICO DE CONDIÇÕES AMBIENTAIS DO TRABALHO (LTCAT) & PPP" if "Previdenciário" in safe_str(process.get("modulo_atuacao")) else "PRÉ-RELATÓRIO DE ANÁLISE PROCESSUAL E PERICIAL"
    _apply_run_style(title_paragraph.add_run(title), size=15, bold=True, color=PRIMARY_COLOR)

    role_paragraph = _new_paragraph(document, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    _apply_run_style(role_paragraph.add_run(f"Módulo: {safe_str(process.get('modulo_atuacao'))}\n"), bold=True)
    _apply_run_style(role_paragraph.add_run(f"Profissional: {safe_str(process.get('papel_profissional') or 'Perito / Consultor')}"), bold=True)

    if "Previdenciário" in safe_str(process.get("modulo_atuacao")):
        _build_previdenciario_document(document, process)
    else:
        _build_judicial_document(document, process)

    if process.get("campo_fotos"):
        _add_title(document, "REGISTROS FOTOGRÁFICOS DE CAMPO", level=3)
        for index, photo in enumerate(process.get("campo_fotos", []), start=1):
            try:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if photo.get("base64"):
                    image_bytes = decode_photo_base64(photo.get("base64"))
                    paragraph.add_run().add_picture(io.BytesIO(image_bytes), width=Inches(4.5))
                elif photo.get("path") and os.path.exists(safe_str(photo.get("path"))):
                    paragraph.add_run().add_picture(safe_str(photo.get("path")), width=Inches(4.5))
                else:
                    warnings.append(f"Foto {index} ignorada: imagem ausente.")
                    continue
            except (ValueError, OSError) as exc:
                warnings.append(f"Foto {index} ignorada: {exc}")
                continue
            legend = safe_str(photo.get("legenda")) or "Registro fotográfico."
            gps = safe_str(photo.get("gps"))
            if gps:
                legend = f"{legend} ({gps})"
            legend_paragraph = document.add_paragraph(f"Figura {index}: {legend}")
            legend_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            legend_paragraph.paragraph_format.line_spacing = 1.5
            for run in legend_paragraph.runs:
                _apply_run_style(run, size=9, italic=True)
            _add_blank_line(document)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer, warnings
