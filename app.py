import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import re
import copy
import io
import logging
import html
import json
import hashlib
from datetime import datetime

from core.config import LOGO_FILE, criar_dados_padrao
from core.database import (
    carregar_dados,
    salvar_processo as _salvar_processo,
    excluir_processo,
    gerar_proximo_id,
    estimar_proximo_id_local,
)
from core.campo_fotos import (
    adicionar_pendentes,
    criar_foto,
    ids_fotos,
    resumo_armazenamento,
    separar_por_limite,
)
from core.photo_storage import enviar_foto, metadados_storage, resolver_bytes_foto
from core.field_records import (
    FIELD_RECORD_FIELDS,
    FIELD_RECORD_KEYS,
    atualizar_campos_memorizados,
    validar_registro_campo,
)
from core.utils import remover_acentos, calcula_altura, comprimir_imagem
from ui import aplicar_estilos
from core.extrajudicial import (
    new_defaults as extrajudicial_defaults, invalidate_review, final_errors,
    parse_docx as parse_extrajudicial_docx, recognized_import,
)
from core.extrajudicial_document import append_extrajudicial, build_extrajudicial_bytes
from ui.extrajudicial import render_extrajudicial

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s [%(name)s] %(message)s",
)
logger = logging.getLogger(__name__)


def salvar_processo(processo_id, dados, novo=False):
    """Grava somente no caso ativo (ou no caso recém-criado), nunca em outro ID."""
    caso_ativo = st.session_state.get("processo_ativo")
    if not novo and processo_id != caso_ativo:
        logger.error("Gravação bloqueada: destino %s difere do caso ativo %s.", processo_id, caso_ativo)
        st.error(
            f"Gravação bloqueada: o destino ({processo_id}) não é o caso ativo ({caso_ativo}). "
            "Nenhum dado foi salvo. Confira o caso selecionado e tente novamente."
        )
        return False
    payload = dados
    if "Previdenciário" in str(dados.get("modulo_atuacao", "")):
        payload = invalidate_review(dados)
    if not _salvar_processo(processo_id, payload):
        return False
    nova_revisao = st.session_state.get("_process_revision_snapshot", {}).get(processo_id)
    if nova_revisao is not None:
        dados["_revision"] = nova_revisao
    return True


def _field_widget_key(field_config: dict, processo_id: str) -> str:
    return f"{field_config['widget_key_prefix']}_{processo_id}"


def _field_placeholder(field_config: dict, processo_id: str) -> str:
    # O ID do caso no placeholder isola rascunho local e ditado por voz por caso.
    return f"{field_config['placeholder']} — caso {processo_id}"


def _descartar_rascunho_campo(processo_id: str) -> None:
    st.session_state.setdefault("campo_rascunhos", {}).pop(processo_id, None)
    st.session_state.pop(f"campo_validation_errors_{processo_id}", None)
    for field in FIELD_RECORD_FIELDS:
        st.session_state.pop(_field_widget_key(field, processo_id), None)


def _descartar_fotos_pendentes(processo_id: str) -> None:
    st.session_state.setdefault("fotos_pendentes", {}).pop(processo_id, None)


def _painel_fotos_key(processo_id: str) -> str:
    return f"fotos_adicionar_aberto_{processo_id}"


def _alternar_painel_fotos(processo_id: str) -> None:
    chave = _painel_fotos_key(processo_id)
    st.session_state[chave] = not st.session_state.get(chave, False)


def _recarregar_caso(processo_id: str) -> None:
    # Na próxima carga, a revisão esperada passa a ser a atual da nuvem.
    st.session_state.get("_process_revision_snapshot", {}).pop(processo_id, None)


def render_caso_ativo(processo_id: str, dados: dict, papel_padrao: str = "Não Definido") -> None:
    partes = [f"Caso Ativo: {processo_id}"]
    nome = str(dados.get("reclamante_nome", "") or "").strip()
    numero = str(dados.get("processo_num", "") or "").strip()
    if nome:
        partes.append(nome)
    if numero:
        partes.append(f"Nº {numero}")
    partes.append(f"Papel: {dados.get('papel_profissional', papel_padrao)}")
    st.info(" | ".join(partes))


def _set_field_widget_value(field_config: dict, processo_id: str, value: str) -> None:
    st.session_state[_field_widget_key(field_config, processo_id)] = value


def _normalizar_dados_documento(dados: dict) -> tuple[dict, list[str]]:
    documento = dict(dados)
    campos_texto = {
        "modulo_atuacao", "papel_profissional", "reclamante_nome", "reclamante_cpf",
        "segurado_nascimento", "profissao_cargo", "reclamada_nome", "reclamada_cnpj",
        "setor", "data_admissao", "relato_inicial", "apr_fisicos", "apr_quimicos",
        "apr_biologicos", "enquadramento_legal_prev", "doc_ltcat", "conclusao_tecnica",
        "extemp_justificativa", "analise_epis_critica", "processo_num", "orgao_julgador",
        "data_autuacao", "valor_causa", "rito_processual", "reclamante_adv",
        "reclamada_adv", "status_contrato", "periodo_imprescrito", "cargos",
        "ultima_remuneracao", "objeto_pericia", "atividades_inicial", "agentes_alegados",
        "pedidos_tecnicos", "preliminares_periciais", "defesa_merito_sst",
        "fase_processual", "campo_data", "campo_horario", "local_diligencia",
        "doc_laudo", "doc_ppp", "doc_pgr", "doc_os", "doc_asos", "doc_outros",
        "quesitos_juizo", "quesitos_autor", "quesitos_reu", "presentes_pericia",
        "campo_declaracoes_autor", "campo_declaracoes_reu", "campo_medicoes",
    }
    for campo in campos_texto:
        valor = documento.get(campo, "")
        if valor is None:
            documento[campo] = ""
        elif isinstance(valor, (str, int, float, bool)):
            documento[campo] = str(valor)
        else:
            return documento, [f"O campo '{campo}' contém um formato incompatível."]

    erros = []
    epis = documento.get("quadro_epis", [])
    if not isinstance(epis, list) or any(not isinstance(item, dict) for item in epis):
        erros.append("A tabela de EPIs possui estrutura inválida.")
    else:
        documento["quadro_epis"] = [
            {campo: str(item.get(campo) or "") for campo in ("descricao", "ca", "data_entrega", "obs")}
            for item in epis
        ]
    fotos = documento.get("campo_fotos", [])
    if not isinstance(fotos, list) or any(
        not isinstance(item, dict) or not isinstance(item.get("base64", ""), str)
        for item in fotos
    ):
        erros.append("A lista de fotos possui estrutura inválida.")
    return documento, erros


def render_field_error(message: str | None) -> None:
    if message:
        st.markdown(
            f"<div class='campo-inline-error' role='alert'>{html.escape(message)}</div>",
            unsafe_allow_html=True,
        )


def render_field_draft_support(processo_id: str) -> None:
    draft_key = f"grahl-campo-draft::{processo_id}"
    field_map = [
        {"key": field["key"], "placeholder": _field_placeholder(field, processo_id)}
        for field in FIELD_RECORD_FIELDS
    ]

    components.html(
        f"""
        <style>
            .campo-draft-status-wrap {{
                display:flex;
                align-items:center;
                justify-content:space-between;
                gap:12px;
                padding:0.6rem 0.8rem;
                border:1px solid #CBD5E1;
                border-radius:10px;
                background:#FFFFFF;
                font-family:'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            }}
            .campo-draft-status {{
                color:#334155;
                font-size:0.9rem;
                font-weight:600;
            }}
            .campo-draft-clear {{
                background:#F8FAFC;
                border:1px solid #CBD5E1;
                border-radius:8px;
                color:#1E293B;
                cursor:pointer;
                font-size:0.85rem;
                font-weight:700;
                min-height:40px;
                padding:0.55rem 0.8rem;
                white-space:nowrap;
            }}
            @media (max-width: 640px) {{
                .campo-draft-status-wrap {{
                    align-items:stretch;
                    flex-direction:column;
                }}
                .campo-draft-clear {{
                    width:100%;
                }}
            }}
        </style>
        <div class="campo-draft-status-wrap">
            <div id="campo_draft_status" class="campo-draft-status">Rascunho somente local: ainda não sincronizado com a nuvem.</div>
            <button id="campo_draft_clear" type="button" class="campo-draft-clear">Limpar rascunho local</button>
        </div>
        <script>
            (function() {{
                const storageKey = {json.dumps(draft_key)};
                const fields = {json.dumps(field_map)};
                const statusEl = document.getElementById("campo_draft_status");
                const clearBtn = document.getElementById("campo_draft_clear");
                const parentDoc = window.parent.document;

                function setStatus(message, color) {{
                    statusEl.textContent = message;
                    statusEl.style.color = color || "#334155";
                }}

                function getField(placeholder) {{
                    const safe = window.CSS && CSS.escape ? CSS.escape(placeholder) : placeholder;
                    return parentDoc.querySelector(
                        `textarea[placeholder="${{safe}}"], input[placeholder="${{safe}}"]`
                    );
                }}

                function setNativeValue(target, value) {{
                    const valueProto = target.tagName === "TEXTAREA"
                        ? window.parent.HTMLTextAreaElement.prototype
                        : window.parent.HTMLInputElement.prototype;
                    const setter = Object.getOwnPropertyDescriptor(valueProto, "value")?.set;
                    if (setter) {{
                        setter.call(target, value);
                    }} else {{
                        target.value = value;
                    }}
                    target.dispatchEvent(new Event("input", {{ bubbles: true }}));
                    target.dispatchEvent(new Event("change", {{ bubbles: true }}));
                }}

                function captureValues() {{
                    const values = {{}};
                    let hasContent = false;
                    fields.forEach((field) => {{
                        const target = getField(field.placeholder);
                        if (!target) {{
                            return;
                        }}
                        values[field.key] = target.value || "";
                        if ((target.value || "").trim()) {{
                            hasContent = true;
                        }}
                    }});
                    return {{ values, hasContent }};
                }}

                function saveDraft() {{
                    const snapshot = captureValues();
                    if (!snapshot.hasContent) {{
                        window.parent.localStorage.removeItem(storageKey);
                        setStatus("Rascunho somente local: ainda não sincronizado com a nuvem.", "#334155");
                        return;
                    }}

                    const payload = {{
                        values: snapshot.values,
                        updatedAt: new Date().toISOString(),
                    }};
                    try {{
                        window.parent.localStorage.setItem(storageKey, JSON.stringify(payload));
                        setStatus("Rascunho somente local salvo; use Salvar para sincronizar com a nuvem.", "#166534");
                    }} catch (error) {{
                        setStatus("Não foi possível salvar o rascunho neste navegador.", "#b91c1c");
                    }}
                }}

                function restoreDraft() {{
                    const raw = window.parent.localStorage.getItem(storageKey);
                    if (!raw) {{
                        return;
                    }}

                    let payload;
                    try {{
                        payload = JSON.parse(raw);
                    }} catch (error) {{
                        window.parent.localStorage.removeItem(storageKey);
                        return;
                    }}
                    const updatedAtMs = Date.parse(payload?.updatedAt || "");
                    if (!Number.isFinite(updatedAtMs) || Date.now() - updatedAtMs > 7 * 24 * 60 * 60 * 1000) {{
                        window.parent.localStorage.removeItem(storageKey);
                        setStatus("Rascunho local expirado (retenção de 7 dias) e removido.", "#b45309");
                        return;
                    }}

                    let restored = 0;
                    fields.forEach((field) => {{
                        const target = getField(field.placeholder);
                        const nextValue = payload?.values?.[field.key] || "";
                        if (!target || !nextValue || (target.value || "").trim()) {{
                            return;
                        }}
                        setNativeValue(target, nextValue);
                        restored += 1;
                    }});

                    if (restored) {{
                        const updatedAt = payload.updatedAt
                            ? new Date(payload.updatedAt).toLocaleString("pt-BR")
                            : "agora";
                        setStatus(`Rascunho local restaurado automaticamente (${{
                            restored
                        }} campo(s), atualizado em ${{updatedAt}}.`, "#0f766e");
                    }}
                }}

                function bindInputs() {{
                    let ready = 0;
                    fields.forEach((field) => {{
                        const target = getField(field.placeholder);
                        if (!target) {{
                            return;
                        }}
                        ready += 1;
                        if (target.dataset.campoDraftBound === "1") {{
                            return;
                        }}
                        target.dataset.campoDraftBound = "1";
                        target.addEventListener("input", saveDraft);
                        target.addEventListener("change", saveDraft);
                    }});
                    return ready;
                }}

                clearBtn.addEventListener("click", function() {{
                    window.parent.localStorage.removeItem(storageKey);
                    setStatus("Rascunho local removido deste navegador.", "#b45309");
                }});

                const setup = window.setInterval(function() {{
                    if (bindInputs() === fields.length) {{
                        restoreDraft();
                        saveDraft();
                        window.clearInterval(setup);
                        window.setInterval(saveDraft, 4000);
                    }}
                }}, 500);
            }})();
        </script>
        """,
        height=88,
    )


def render_voice_dictation_control(widget_placeholder: str, field_label: str, control_key: str) -> None:
    """Renderiza botão de ditado por voz para um campo Streamlit identificado por placeholder."""
    safe_key = re.sub(r"[^a-zA-Z0-9_-]", "_", control_key)
    btn_id = f"voice_btn_{safe_key}"
    status_id = f"voice_status_{safe_key}"

    components.html(
        f"""
        <div style="display:flex; align-items:center; gap:8px; flex-wrap:wrap; margin:0.15rem 0 0.8rem 0;">
            <button id="{btn_id}" type="button" aria-label="Ditado por voz para {html.escape(field_label)}"
                style="border:1px solid #cbd5e1; border-radius:8px; padding:8px 12px; min-height:44px; cursor:pointer; background:#ffffff; font-size:14px;">
                🎤 Ditar
            </button>
            <span id="{status_id}" style="font-size:12px; color:#334155;">Toque para ditar.</span>
        </div>
        <style>
            @media (max-width: 640px) {{
                #{btn_id} {{
                    width: 100%;
                    justify-content: center;
                }}
                #{status_id} {{
                    width: 100%;
                }}
            }}
        </style>
        <script>
            (function() {{
                const btn = document.getElementById({json.dumps(btn_id)});
                const status = document.getElementById({json.dumps(status_id)});
                const targetPlaceholder = {json.dumps(widget_placeholder)};
                const SpeechRecognition = window.parent.SpeechRecognition || window.parent.webkitSpeechRecognition;

                function setStatus(msg, color) {{
                    status.textContent = msg;
                    status.style.color = color || "#334155";
                }}

                function applyTranscript(transcript) {{
                    const safePlaceholder = window.CSS && CSS.escape ? CSS.escape(targetPlaceholder) : targetPlaceholder;
                    const target = window.parent.document.querySelector(
                        `textarea[placeholder="${{safePlaceholder}}"], input[placeholder="${{safePlaceholder}}"]`
                    );
                    if (!target) {{
                        setStatus("Não foi possível localizar o campo.", "#b91c1c");
                        return;
                    }}

                    const currentValue = target.value || "";
                    const hasText = currentValue.trim().length > 0;
                    const normalized = transcript.trim();
                    if (!normalized) {{
                        setStatus("Nenhum áudio reconhecido.", "#b45309");
                        return;
                    }}
                    const prefix = hasText && !/[\\s\\n]$/.test(currentValue) ? " " : "";
                    const insertText = prefix + normalized;
                    const start = Number.isInteger(target.selectionStart) ? target.selectionStart : currentValue.length;
                    const end = Number.isInteger(target.selectionEnd) ? target.selectionEnd : currentValue.length;
                    const nextValue = currentValue.slice(0, start) + insertText + currentValue.slice(end);

                    const valueProto = target.tagName === "TEXTAREA"
                        ? window.parent.HTMLTextAreaElement.prototype
                        : window.parent.HTMLInputElement.prototype;
                    const setter = Object.getOwnPropertyDescriptor(valueProto, "value")?.set;
                    if (setter) {{
                        setter.call(target, nextValue);
                    }} else {{
                        target.value = nextValue;
                    }}

                    const nextCursor = start + insertText.length;
                    target.selectionStart = nextCursor;
                    target.selectionEnd = nextCursor;
                    target.dispatchEvent(new Event("input", {{ bubbles: true }}));
                    target.dispatchEvent(new Event("change", {{ bubbles: true }}));
                    setStatus("Texto transcrito adicionado ao campo.", "#166534");
                }}

                if (!SpeechRecognition) {{
                    btn.disabled = true;
                    btn.style.opacity = "0.6";
                    btn.style.cursor = "not-allowed";
                    setStatus("Ditado por voz não suportado neste navegador.", "#475569");
                    return;
                }}

                btn.addEventListener("click", function() {{
                    const recognition = new SpeechRecognition();
                    recognition.lang = "pt-BR";
                    recognition.interimResults = false;
                    recognition.maxAlternatives = 1;
                    recognition.continuous = false;

                    btn.disabled = true;
                    setStatus("🎙️ Ouvindo... fale agora.", "#0f766e");

                    recognition.onresult = function(event) {{
                        const transcript = event.results?.[0]?.[0]?.transcript || "";
                        applyTranscript(transcript);
                    }};

                    recognition.onerror = function(event) {{
                        const messageMap = {{
                            "not-allowed": "Permissão de microfone negada.",
                            "service-not-allowed": "Serviço de voz indisponível neste dispositivo.",
                            "audio-capture": "Microfone não encontrado.",
                            "network": "Falha de rede ao transcrever.",
                            "no-speech": "Nenhuma fala detectada."
                        }};
                        setStatus(messageMap[event.error] || `Erro de voz: ${{event.error}}.`, "#b91c1c");
                    }};

                    recognition.onend = function() {{
                        btn.disabled = false;
                    }};

                    try {{
                        recognition.start();
                    }} catch (err) {{
                        btn.disabled = false;
                        setStatus("Não foi possível iniciar o ditado agora.", "#b91c1c");
                    }}
                }});
            }})();
        </script>
        """,
        height=54,
    )


# === MOTOR DE EXTRAÇÃO DUPLO (TAGS + NATURAL) ===
def parse_pre_relatorio(doc):
    dados = {}
    texto_completo = ""
    linhas = []
    
    for p in doc.paragraphs:
        texto_completo += p.text + "\n"
        if p.text.strip(): linhas.append(p.text.strip())
        
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join([c.text.strip() for c in row.cells if c.text.strip()])
            if row_text:
                texto_completo += row_text + "\n"
                linhas.append(row_text)

    # 1. MOTOR DE TAGS DA IA (Prioridade Máxima)
    if "[START_" in texto_completo:
        def extrair_bloco(nome_tag):
            padrao = r"\[START_" + nome_tag + r"\](.*?)\[END_" + nome_tag + r"\]"
            match = re.search(padrao, texto_completo, re.DOTALL | re.IGNORECASE)
            if match:
                bloco = match.group(1).strip()
                if "[informação" in bloco.lower() or "não consta" in bloco.lower() or bloco.lower() in ["n/a", "na", "-"]:
                    return ""
                return bloco
            return ""

        dados["modulo_atuacao"] = extrair_bloco("MODULO")
        papel = extrair_bloco("PAPEL").lower()
        if "extrajudicial" in papel or "consultor" in papel: dados["papel_profissional"] = "Perito Extrajudicial / Consultor"
        elif "reclamante" in papel or "autor" in papel: dados["papel_profissional"] = "Assistente Técnico da Reclamante"
        elif "assistente" in papel or "reclamada" in papel: dados["papel_profissional"] = "Assistente Técnico da Reclamada"
        elif papel: dados["papel_profissional"] = "Perito do Juízo"
        
        dados["processo_num"] = extrair_bloco("PROCESSO_NUM")
        dados["orgao_julgador"] = extrair_bloco("ORGAO_JULGADOR")
        dados["data_autuacao"] = extrair_bloco("DATA_AUTUACAO")
        dados["valor_causa"] = extrair_bloco("VALOR_CAUSA")
        dados["rito_processual"] = extrair_bloco("RITO")
        
        dados["reclamante_nome"] = extrair_bloco("RECLAMANTE_NOME")
        dados["reclamante_cpf"] = extrair_bloco("RECLAMANTE_CPF")
        dados["reclamante_adv"] = extrair_bloco("RECLAMANTE_ADV")
        dados["segurado_nascimento"] = extrair_bloco("NASCIMENTO")
        
        dados["reclamada_nome"] = extrair_bloco("RECLAMADA_NOME")
        dados["reclamada_cnpj"] = extrair_bloco("RECLAMADA_CNPJ")
        dados["reclamada_adv"] = extrair_bloco("RECLAMADA_ADV")
        
        dados["data_admissao"] = extrair_bloco("DATA_ADMISSAO")
        dados["status_contrato"] = extrair_bloco("STATUS_CONTRATO")
        dados["periodo_imprescrito"] = extrair_bloco("PERIODO_IMPRESCRITO")
        
        cargos = extrair_bloco("CARGOS")
        dados["cargos"] = cargos
        dados["profissao_cargo"] = cargos
        
        dados["setor"] = extrair_bloco("SETOR")
        dados["ultima_remuneracao"] = extrair_bloco("ULTIMA_REMUNERACAO")
        dados["objeto_pericia"] = extrair_bloco("OBJETO_PERICIA")
        
        relato = extrair_bloco("RELATO_INICIAL")
        dados["atividades_inicial"] = relato
        dados["relato_inicial"] = relato
        
        dados["agentes_alegados"] = extrair_bloco("AGENTES_ALEGADOS")
        dados["apr_fisicos"] = extrair_bloco("APR_FISICOS")
        dados["apr_quimicos"] = extrair_bloco("APR_QUIMICOS")
        dados["apr_biologicos"] = extrair_bloco("APR_BIOLOGICOS")
        dados["enquadramento_legal_prev"] = extrair_bloco("ENQUADRAMENTO")
        
        dados["pedidos_tecnicos"] = extrair_bloco("PEDIDOS_TECNICOS")
        dados["preliminares_periciais"] = extrair_bloco("PRELIMINARES")
        dados["defesa_merito_sst"] = extrair_bloco("DEFESA_MERITO")
        
        dados["fase_processual"] = extrair_bloco("FASE_PROCESSUAL")
        dados["campo_data"] = extrair_bloco("DATA_VISTORIA")
        dados["campo_horario"] = extrair_bloco("HORARIO_VISTORIA")
        dados["local_diligencia"] = extrair_bloco("LOCAL_VISTORIA")
        
        dados["doc_ltcat"] = extrair_bloco("DOC_LTCAT")
        dados["doc_laudo"] = extrair_bloco("DOC_LAUDO")
        dados["doc_ppp"] = extrair_bloco("DOC_PPP")
        dados["doc_pgr"] = extrair_bloco("DOC_PGR")
        dados["doc_os"] = extrair_bloco("DOC_OS")
        dados["doc_asos"] = extrair_bloco("DOC_ASOS")
        dados["doc_outros"] = extrair_bloco("DOC_OUTROS")
        dados["analise_epis_critica"] = extrair_bloco("SINTESE_EPIS")
        
        dados["quesitos_juizo"] = extrair_bloco("QUESITOS_JUIZO")
        dados["quesitos_autor"] = extrair_bloco("QUESITOS_AUTOR")
        dados["quesitos_reu"] = extrair_bloco("QUESITOS_REU")

        quadro_epis_texto = extrair_bloco("QUADRO_EPIS")
        epis_extraidos = []
        if quadro_epis_texto:
            linhas_epi = quadro_epis_texto.strip().split('\n')
            for linha in linhas_epi:
                if '|' in linha and '---' not in linha and 'descrição' not in linha.lower():
                    parts = [p.strip() for p in linha.split('|')]
                    if len(parts) >= 2:
                        desc = parts[0]
                        ca = parts[1]
                        data = parts[2] if len(parts) > 2 else ""
                        obs = parts[3] if len(parts) > 3 else ""
                        if desc and "[informação" not in desc.lower() and "n/a" not in desc.lower() and desc != "-":
                            epis_extraidos.append({"descricao": desc, "ca": ca, "data_entrega": data, "obs": obs})
        dados["quadro_epis"] = epis_extraidos
        
        return dados

    # 2. MOTOR DE TEXTO NATURAL (Fallback para documentos sem tags)
    texto = "\n".join(linhas).replace("**", "").replace("*", "")

    def get_single(padrao):
        m = re.search(r"(?:" + padrao + r")[\s]*[:|][\s]*([^\n]+)", texto, re.IGNORECASE)
        if m:
            v = m.group(1).replace("_", "").strip()
            if v and "[informação" not in v.lower(): return v
        return ""

    def get_multi(padrao, next_sections):
        lookahead = r"(?=\n\d+\.\s*|(?:" + "|".join(next_sections) + r")|$)"
        m = re.search(r"(?:" + padrao + r")[\s]*[:|]?[\s]*(.*?)" + lookahead, texto, re.IGNORECASE | re.DOTALL)
        if m:
            v = m.group(1).replace("_", "").strip()
            if v and "[informação" not in v.lower(): return v
        return ""

    dados["processo_num"] = get_single(r"Número do Processo[^\n:]*|Processo")
    dados["orgao_julgador"] = get_single(r"Órgão Julgador[^\n:]*|Vara")
    dados["data_autuacao"] = get_single(r"Data de Autuação[^\n:]*|Ajuizamento.*?")
    dados["valor_causa"] = get_single(r"Valor da Causa")
    dados["rito_processual"] = get_single(r"Rito Processual")

    rec = get_single(r"Reclamante[^\n:]*|Nome do Segurado[^\n:]*|Nome[^\n:]*")
    if rec:
        m_cpf = re.search(r"\((?:CPF|NIT|PIS)\s*([\d\.\-\/]+)\)", rec, re.IGNORECASE)
        if m_cpf:
            dados["reclamante_cpf"] = m_cpf.group(1)
            rec = re.sub(r"\((?:CPF|NIT|PIS).*?\)", "", rec, flags=re.IGNORECASE).strip()
        dados["reclamante_nome"] = rec
    if not dados.get("reclamante_cpf"): dados["reclamante_cpf"] = get_single(r"CPF[^\n:]*|NIT[^\n:]*|PIS[^\n:]*")

    recd = get_single(r"Reclamada[^\n:]*|Razão Social[^\n:]*")
    if recd:
        m_cnpj = re.search(r"CNPJ[\s:]*([\d\.\-\/]+)", recd, re.IGNORECASE)
        if m_cnpj:
            dados["reclamada_cnpj"] = m_cnpj.group(1)
            recd = re.sub(r"\(?CNPJ.*?\)?", "", recd, flags=re.IGNORECASE).strip()
        dados["reclamada_nome"] = recd.strip(" ()-")
    if not dados.get("reclamada_cnpj"): dados["reclamada_cnpj"] = get_single(r"CNPJ[^\n:]*")

    advs = re.findall(r"Advogados[\s]*[:|][\s]*([^\n]+)", texto, re.IGNORECASE)
    if len(advs) >= 1: dados["reclamante_adv"] = advs[0].strip()
    if len(advs) >= 2: dados["reclamada_adv"] = advs[1].strip()

    dados["data_admissao"] = get_single(r"Data de Admissão[^\n:]*|Período Avaliado")
    dados["status_contrato"] = get_single(r"Status do Contrato[^\n:]*")
    dados["periodo_imprescrito"] = get_single(r"Período Imprescrito[^\n:]*")
    
    cargo = get_single(r"Cargo[^\n:]*|Profissão[^\n:]*")
    if cargo:
        dados["cargos"] = cargo
        dados["profissao_cargo"] = cargo
        
    dados["segurado_nascimento"] = get_single(r"Data de Nascimento[^\n:]*")
    dados["setor"] = get_single(r"Setor[^\n:]*|Lotação[^\n:]*")
    dados["ultima_remuneracao"] = get_single(r"Última Remuneração[^\n:]*")
    dados["objeto_pericia"] = get_single(r"Objeto da Perícia[^\n:]*|Objeto de Análise[^\n:]*|Objetivo[^\n:]*")

    paradas = [r"Agentes Nocivos", r"Agentes Físicos", r"Pedidos Técnicos", r"Preliminares Periciais", r"Defesa de Mérito", r"Fase Processual", r"Data da Vistoria", r"Local / Endereço", r"Análise do LTCAT", r"Laudo de Insalubridade", r"Análise do PPP", r"Análise do PGR", r"Ordens de Serviço", r"ASOs / PCMSO", r"Outros Documentos", r"Síntese e Análise", r"9\.1", r"9\.2", r"9\.3", r"Pessoas Presentes", r"Informações prestadas", r"Medições Realizadas"]
    
    dados["atividades_inicial"] = get_multi(r"Atividades Descritas[^\n:]*|Atividades Típicas[^\n:]*|Relato das Atividades[^\n:]*|Relato Inicial[^\n:]*", paradas)
    if not dados.get("relato_inicial"): dados["relato_inicial"] = dados.get("atividades_inicial", "")

    dados["agentes_alegados"] = get_multi(r"Agentes Nocivos[^\n:]*|Riscos Alegados[^\n:]*", paradas)
    dados["apr_fisicos"] = get_multi(r"Agentes Físicos[^\n:]*", paradas)
    dados["apr_quimicos"] = get_multi(r"Agentes Químicos[^\n:]*", paradas)
    dados["apr_biologicos"] = get_multi(r"Agentes Biológicos[^\n:]*", paradas)
    dados["enquadramento_legal_prev"] = get_multi(r"Enquadramento Legal[^\n:]*", paradas)
    
    dados["pedidos_tecnicos"] = get_multi(r"Pedidos Técnicos[^\n:]*", paradas)
    dados["preliminares_periciais"] = get_multi(r"Preliminares Periciais[^\n:]*", paradas)
    dados["defesa_merito_sst"] = get_multi(r"Defesa de Mérito[^\n:]*", paradas)

    dados["fase_processual"] = get_multi(r"Fase Processual Atual[^\n:]*", paradas)
    dados["campo_data"] = get_multi(r"Data da Vistoria[^\n:]*", paradas)
    dados["campo_horario"] = get_multi(r"Horário[^\n:]*", paradas)
    dados["local_diligencia"] = get_multi(r"Local[^\n:]*|Endereço[^\n:]*", paradas)

    dados["doc_ltcat"] = get_multi(r"LTCAT[^\n:]*", paradas)
    dados["doc_laudo"] = get_multi(r"LAUDO DE INSALUBRIDADE[^\n:]*|Laudos Prévios[^\n:]*", paradas)
    dados["doc_ppp"] = get_multi(r"PPP[^\n:]*", paradas)
    dados["doc_pgr"] = get_multi(r"PGR[^\n:]*|PPRA[^\n:]*", paradas)
    dados["doc_os"] = get_multi(r"ORDENS DE SERVIÇO[^\n:]*", paradas)
    dados["doc_asos"] = get_multi(r"ASOs[^\n:]*|PCMSO[^\n:]*", paradas)
    dados["doc_outros"] = get_multi(r"Outros Documentos[^\n:]*", paradas)

    dados["analise_epis_critica"] = get_multi(r"Síntese e Análise Crítica de EPIs[^\n:]*", paradas)

    dados["quesitos_juizo"] = get_multi(r"9\.1\.\s+Quesitos do Juízo", [r"\n9\.2\."])
    dados["quesitos_autor"] = get_multi(r"9\.2\.\s+Quesitos do Reclamante.*?", [r"\n9\.3\."])
    dados["quesitos_reu"] = get_multi(r"9\.3\.\s+Quesitos da Reclamada.*?", [r"\n10\.\s+LEVANTAMENTOS", r"\nPessoas Presentes"])

    epis_extraidos = []
    for table in doc.tables:
        if len(table.rows) > 0:
            hdr = [cell.text.lower() for cell in table.rows[0].cells]
            if any("descri" in h or "epi" in h for h in hdr):
                for row in table.rows[1:]:
                    cells = row.cells
                    desc = cells[0].text.strip() if len(cells) > 0 else ""
                    ca = cells[1].text.strip() if len(cells) > 1 else ""
                    data = cells[2].text.strip() if len(cells) > 2 else ""
                    obs = cells[3].text.strip() if len(cells) > 3 else ""
                    desc = desc.replace("**", "").replace("*", "")
                    if desc and "[informação" not in desc.lower() and desc != "-":
                        epis_extraidos.append({"descricao": desc, "ca": ca, "data_entrega": data, "obs": obs})
    if epis_extraidos and not dados.get("quadro_epis"): 
        dados["quadro_epis"] = epis_extraidos

    return dados

icon_config = LOGO_FILE if os.path.exists(LOGO_FILE) else "🛡️"
st.set_page_config(page_title="GRAHL Perícias", page_icon=icon_config, layout="wide")

if "processo_ativo" not in st.session_state: st.session_state.processo_ativo = None
if "menu_opcao" not in st.session_state: st.session_state.menu_opcao = "➕ Novo Processo / Caso"
if "parsed_data" not in st.session_state: st.session_state.parsed_data = {}
if "gps_por_caso" not in st.session_state: st.session_state.gps_por_caso = {}
if "campo_rascunhos" not in st.session_state: st.session_state.campo_rascunhos = {}
if "fotos_pendentes" not in st.session_state: st.session_state.fotos_pendentes = {}
if "confirmar_exclusao_target" not in st.session_state: st.session_state.confirmar_exclusao_target = ""
if "confirmar_exclusao_input_nonce" not in st.session_state: st.session_state.confirmar_exclusao_input_nonce = 0
if "uploader_key" not in st.session_state: st.session_state.uploader_key = 0  
if "campo_valores_memorizados" not in st.session_state: st.session_state.campo_valores_memorizados = {}

aplicar_estilos()

db_processos = carregar_dados()
if st.session_state.get("firebase_load_error"):
    st.error(
        "Os dados existentes não foram carregados. Não crie nem edite casos até que a conexão "
        "com o banco seja restaurada; nenhuma cópia local substitui os dados da nuvem."
    )

col_logo, col_titulo = st.columns([1, 6])
with col_logo:
    logo_encontrado = LOGO_FILE if os.path.exists(LOGO_FILE) else None
    if logo_encontrado: st.image(logo_encontrado, width=140)
    else: st.markdown("### 🛡️ **GRAHL**")

with col_titulo:
    st.markdown("<h1 style='margin:0; font-size: 1.7rem;'>GRAHL Perícias</h1>", unsafe_allow_html=True)
    st.markdown("<p style='color: #64748B; margin:0; font-size: 1rem; font-weight: 500;'>Assistência Técnica, Perícias Judiciais e Laudos Previdenciários.</p>", unsafe_allow_html=True)

st.markdown("<hr style='margin:0.5rem 0 1.5rem 0; border: none; height: 1px; background-color: #CBD5E1;'>", unsafe_allow_html=True)

def trocar_menu(acao):
    st.session_state.menu_opcao = acao


def salvar_processo_com_feedback(id_proc, dados_proc, sucesso_msg, erro_msg="❌ Não foi possível salvar os dados. Tente novamente."):
    if salvar_processo(id_proc, dados_proc):
        if sucesso_msg:
            st.toast(sucesso_msg, icon="💾")
        return True
    st.error(erro_msg)
    return False

st.sidebar.markdown("<h2 style='color: #FFFFFF; font-size: 1.3rem; margin-bottom: 1rem;'>📁 Painel de Controle</h2>", unsafe_allow_html=True)
acoes_menu = ["➕ Novo Processo / Caso", "✏️ Dados, Escritório & SST", "🚜 Diligência de Campo & Fotos", "🗑️ Excluir Processo", "📄 Gerar Documento Word Final"]
for acao in acoes_menu: st.sidebar.button(acao, on_click=trocar_menu, args=(acao,), width="stretch")

opcao = st.session_state.menu_opcao

st.sidebar.markdown("<hr style='border: none; height: 1px; background-color: rgba(255,255,255,0.2); margin: 1.5rem 0;'>", unsafe_allow_html=True)
st.sidebar.markdown("<p style='color: #FFFFFF; font-weight: 600; font-size: 14px; margin-bottom: 0.3rem;'>🔎 Pesquisa Rápida:</p>", unsafe_allow_html=True)

termo_busca_geral = st.sidebar.text_input("Busca", placeholder="ID, Processo, Nome ou Empresa...", label_visibility="collapsed").strip()
processos_filtrados = []
termo_limpo = remover_acentos(termo_busca_geral)

if db_processos:
    for k, v in db_processos.items():
        if not termo_limpo:
            processos_filtrados.append(k)
        else:
            id_norm = remover_acentos(k)
            proc_num_norm = remover_acentos(str(v.get("processo_num", "")))
            reclamada_norm = remover_acentos(str(v.get("reclamada_nome", "")))
            reclamante_nome = remover_acentos(str(v.get("reclamante_nome", "")))
            if (termo_limpo in id_norm or termo_limpo in proc_num_norm or termo_limpo in reclamada_norm or termo_limpo in reclamante_nome):
                processos_filtrados.append(k)

st.sidebar.markdown("<br>", unsafe_allow_html=True)

SELETOR_CASO_KEY = "seletor_caso_id"


def _rotulo_caso(k):
    v = db_processos.get(k, {})
    label = f"[{k}] {v.get('reclamante_nome', 'Sem Nome')}"
    empresa = v.get("reclamada_nome", "")
    num_p = v.get("processo_num", "")
    if empresa: label += f" x {empresa}"
    elif num_p: label += f" ({num_p})"
    return label


# Recarga/retorno ao app abre uma nova sessão: restaura o caso escolhido pela URL.
if st.session_state.processo_ativo is None and st.query_params.get("caso") in db_processos:
    st.session_state.processo_ativo = st.query_params.get("caso")
if st.session_state.processo_ativo not in db_processos:
    st.session_state.processo_ativo = (processos_filtrados or list(db_processos) or [None])[0]

if st.session_state.processo_ativo is not None:
    # A pesquisa só filtra a lista; nunca troca o caso ativo sem ação explícita.
    opcoes_casos = list(processos_filtrados)
    if st.session_state.processo_ativo not in opcoes_casos:
        opcoes_casos.insert(0, st.session_state.processo_ativo)
        st.sidebar.caption("Caso ativo mantido; ele não corresponde à pesquisa.")

    def atualizar_processo():
        novo_id = st.session_state[SELETOR_CASO_KEY]
        if novo_id != st.session_state.processo_ativo:
            st.session_state.processo_ativo = novo_id
            st.session_state.menu_opcao = "✏️ Dados, Escritório & SST"

    # Opções são IDs estáveis; o estado do widget é sincronizado antes de instanciá-lo
    # para que a caixa sempre exiba o caso realmente ativo (ex.: após criar/excluir).
    st.session_state[SELETOR_CASO_KEY] = st.session_state.processo_ativo
    st.sidebar.selectbox(
        "Processo / Caso Selecionado:", options=opcoes_casos, format_func=_rotulo_caso,
        key=SELETOR_CASO_KEY, on_change=atualizar_processo,
    )
    processo_id_selecionado = st.session_state.processo_ativo
    if st.query_params.get("caso") != processo_id_selecionado:
        st.query_params["caso"] = processo_id_selecionado
else:
    st.sidebar.warning("⚠️ Nenhum caso encontrado.")
    processo_id_selecionado = "Nenhum caso cadastrado"
    if "caso" in st.query_params:
        del st.query_params["caso"]

aviso_pendencias = st.sidebar.empty()

if opcao == "➕ Novo Processo / Caso":
    st.markdown("### Cadastrar Novo Caso ou Processo")
    proximo_id = estimar_proximo_id_local(db_processos)
    st.info(f"✨ O próximo ID gerado automaticamente é: **{proximo_id}**")

    st.markdown("<br>", unsafe_allow_html=True)
    modulo_escolhido = st.radio("Selecione o Módulo de Atuação:", [
        "⚖️ Perícia Judicial Trabalhista (SST / Insalubridade / Periculosidade / Aposentadoria Especial)", 
        "📄 Laudo Extrajudicial Previdenciário (LTCAT + PPP Extemporâneo/Contemporâneo)"
    ])

    is_prev_mod = "Previdenciário" in modulo_escolhido

    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("#### 📥 Importar Documento Base (.docx) — opcional")
    import_label = ("Selecione o arquivo Word (campos reconhecidos; importação manual)."
                    if is_prev_mod else "Selecione o arquivo Word (preenchimento automático).")
    arquivo_importado = st.file_uploader(import_label, type=["docx"], key=f"uploader_{st.session_state.uploader_key}")
    import_signature = (modulo_escolhido,
                        hashlib.sha256(arquivo_importado.getvalue()).hexdigest()
                        if arquivo_importado is not None else None)
    if st.session_state.get("novo_import_signature") != import_signature:
        st.session_state.parsed_data = {}
        st.session_state.novo_import_signature = import_signature
    
    if arquivo_importado is not None:
        try:
            doc_ext = Document(arquivo_importado)
            dados_extraidos = (parse_extrajudicial_docx(doc_ext, parse_pre_relatorio)
                               if is_prev_mod else parse_pre_relatorio(doc_ext))
            confirmar_importacao = False
            if is_prev_mod:
                st.json(dados_extraidos)
                st.caption("Prévia somente. Nenhum campo digitado será substituído sem confirmação.")
                confirmar_importacao = st.button(
                    "Usar campos reconhecidos no novo caso",
                    key=f"confirmar_import_novo_{st.session_state.uploader_key}")
            if not is_prev_mod or confirmar_importacao:
                st.session_state.parsed_data = dados_extraidos
                for field, short in (("processo_num", "num"), ("reclamante_nome", "nome"),
                                     ("reclamada_nome", "emp")):
                    if field in dados_extraidos:
                        st.session_state[f"novo_{short}_{st.session_state.uploader_key}"] = dados_extraidos[field]
                if is_prev_mod:
                    st.success("Importação confirmada. Revise os dados antes de criar o caso.")
        except Exception as e:
            st.session_state.parsed_data = {}
            st.error(f"Erro ao ler o arquivo: {e}")
    elif st.session_state.get("parsed_data"):
        st.session_state.parsed_data = {}

    st.markdown("<br>", unsafe_allow_html=True)
    with st.form("form_novo"):
        parsed = st.session_state.get("parsed_data", {})
        if is_prev_mod:
            lbl_num = "Identificação / Protocolo ou NIT/CPF"
            lbl_nome = "Nome do Segurado"
            lbl_emp = "Empresa / Tomador / Empregador"
            def_num = parsed.get("processo_num", "Extrajudicial Previdenciário")
        else:
            lbl_num = "Número do Processo (CNJ)"
            lbl_nome = "Reclamante (Autor/Autora)"
            lbl_emp = "Reclamada (Ré / Empresa)"
            def_num = parsed.get("processo_num", "")

        p_num = st.text_input(lbl_num, value=def_num, key=f"novo_num_{st.session_state.uploader_key}")
        p_nome = st.text_input(lbl_nome, value=parsed.get("reclamante_nome", ""), key=f"novo_nome_{st.session_state.uploader_key}")
        p_empresa = st.text_input(lbl_emp, value=parsed.get("reclamada_nome", ""), key=f"novo_emp_{st.session_state.uploader_key}")

        st.markdown("<br>", unsafe_allow_html=True)
        btn_criar_caso = st.form_submit_button("Criar Caso Completo", type="primary", width="stretch")

        if btn_criar_caso:
            id_novo_processo = gerar_proximo_id(db_processos)
            p_novo = criar_dados_padrao()
            if is_prev_mod:
                p_novo.update(extrajudicial_defaults())
                if parsed:
                    parsed = recognized_import(parsed)
            if parsed:
                for k, v in parsed.items():
                    if is_prev_mod or v: p_novo[k] = v
            if is_prev_mod:
                p_novo["modulo_atuacao"] = "Laudo Extrajudicial Previdenciário (LTCAT+PPP)"
                p_novo["reclamante_nome"] = p_nome
                p_novo["reclamada_nome"] = p_empresa
                p_novo["processo_num"] = p_num
            else:
                p_novo["modulo_atuacao"] = "Perícia Judicial Trabalhista"
                p_novo["reclamante_nome"] = p_nome
                p_novo["reclamada_nome"] = p_empresa
                p_novo["processo_num"] = p_num
            
            if id_novo_processo in db_processos:
                st.error(f"❌ O ID {id_novo_processo} já existe; nenhum caso foi sobrescrito. Tente novamente.")
            elif salvar_processo(id_novo_processo, p_novo, novo=True):
                st.session_state.parsed_data = {}
                st.session_state.uploader_key += 1
                st.session_state.processo_ativo = id_novo_processo
                st.session_state.menu_opcao = "✏️ Dados, Escritório & SST"
                st.toast(f"✅ Caso {id_novo_processo} criado com sucesso!", icon="💾")
                st.rerun()
            else:
                st.error("❌ Não foi possível criar o novo caso na nuvem.")

elif opcao == "✏️ Dados, Escritório & SST":
    if not db_processos or processo_id_selecionado == "Nenhum caso cadastrado":
        st.warning("Cadastre ou selecione um caso no menu lateral para começar.")
    else:
        p_atual = db_processos[processo_id_selecionado]
        is_prev = "Previdenciário" in p_atual.get("modulo_atuacao", "")
        
        render_caso_ativo(processo_id_selecionado, p_atual)
        st.caption(
            "Cada aba grava somente ao clicar no respectivo botão Salvar. Alterações não salvas "
            "são descartadas ao trocar de caso e nunca são gravadas em outro caso."
        )

        if is_prev:
            render_extrajudicial(
                processo_id_selecionado, p_atual,
                lambda case_id, data: salvar_processo_com_feedback(case_id, data, "✅ Dados salvos!"),
                parse_pre_relatorio,
            )
        else:
            tab1, tab2, tab3, tab4, tab5 = st.tabs(["1️⃣ Identificação & Partes", "2️⃣ Contrato & Sínteses", "3️⃣ SST & Documentos", "4️⃣ Planilha de EPIs", "5️⃣ Quesitos"])
            with tab1:
                st.markdown("### 1. Papel Profissional, Tipos de Perícia & Identificação")
                with st.form(f"ft_1_{processo_id_selecionado}"):
                    col_p1, col_p2 = st.columns(2)
                    with col_p1:
                        roles = ["Perito do Juízo", "Assistente Técnico da Reclamante", "Assistente Técnico da Reclamada", "Perito Extrajudicial / Consultor"]
                        atual_role = p_atual.get("papel_profissional", "Perito do Juízo")
                        if atual_role not in roles: roles.append(atual_role)
                        p_atual["papel_profissional"] = st.selectbox("Meu Papel no Processo", roles, index=roles.index(atual_role))
                    with col_p2:
                        p_atual["tipos_pericia"] = st.multiselect("Tipos de Perícia Envolvidos", ["Insalubridade (NR-15)", "Periculosidade (NR-16)", "Aposentadoria Especial (PPP/LTCAT)"], default=p_atual.get("tipos_pericia", ["Insalubridade (NR-15)"]))

                    st.markdown("<br>", unsafe_allow_html=True)
                    col1, col2 = st.columns(2)
                    with col1:
                        p_atual["processo_num"] = st.text_input("Número do Processo (CNJ)", value=p_atual.get("processo_num", ""))
                        p_atual["orgao_julgador"] = st.text_input("Órgão Julgador / Vara", value=p_atual.get("orgao_julgador", ""))
                        p_atual["data_autuacao"] = st.text_input("Data de Autuação (Ajuizamento)", value=p_atual.get("data_autuacao", ""))
                        p_atual["valor_causa"] = st.text_input("Valor da Causa", value=p_atual.get("valor_causa", ""))
                        p_atual["rito_processual"] = st.text_input("Rito Processual", value=p_atual.get("rito_processual", ""))
                    with col2:
                        p_atual["reclamante_nome"] = st.text_input("Reclamante (Autor/Autora)", value=p_atual.get("reclamante_nome", ""))
                        p_atual["reclamante_cpf"] = st.text_input("CPF do Reclamante", value=p_atual.get("reclamante_cpf", ""))
                        p_atual["reclamante_adv"] = st.text_input("Advogados do Reclamante (Nomes e OAB)", value=p_atual.get("reclamante_adv", ""))
                        p_atual["reclamada_nome"] = st.text_input("Reclamada (Ré / Empresa)", value=p_atual.get("reclamada_nome", ""))
                        p_atual["reclamada_cnpj"] = st.text_input("CNPJ da Reclamada", value=p_atual.get("reclamada_cnpj", ""))
                        p_atual["reclamada_adv"] = st.text_input("Advogados da Reclamada (Nomes e OAB)", value=p_atual.get("reclamada_adv", ""))

                    if st.form_submit_button("💾 Salvar Identificação"):
                        salvar_processo_com_feedback(
                            processo_id_selecionado,
                            p_atual,
                            "✅ Salvo!",
                        )

            with tab2:
                st.markdown("### 2. Contrato & Sínteses")
                with st.form(f"ft_2_{processo_id_selecionado}"):
                    col3, col4 = st.columns(2)
                    with col3:
                        p_atual["data_admissao"] = st.text_input("Data de Admissão", value=p_atual.get("data_admissao", ""))
                        p_atual["status_contrato"] = st.text_input("Status do Contrato (Demissão / Ativo)", value=p_atual.get("status_contrato", ""))
                        p_atual["periodo_imprescrito"] = st.text_input("Período Imprescrito (Alvo da Perícia)", value=p_atual.get("periodo_imprescrito", ""))
                        p_atual["cargos"] = st.text_input("Cargo(s) / Função(ões)", value=p_atual.get("cargos", ""))
                    with col4:
                        p_atual["setor"] = st.text_input("Setor / Lotação / Local", value=p_atual.get("setor", ""))
                        p_atual["ultima_remuneracao"] = st.text_input("Última Remuneração", value=p_atual.get("ultima_remuneracao", ""))

                    p_atual["objeto_pericia"] = st.text_input("Objeto de Análise / Perícia", value=p_atual.get("objeto_pericia", ""))
                    v_ati = p_atual.get("atividades_inicial", "")
                    p_atual["atividades_inicial"] = st.text_area("Atividades Descritas na Inicial", value=v_ati, height=calcula_altura(v_ati, 100))
                    v_age = p_atual.get("agentes_alegados", "")
                    p_atual["agentes_alegados"] = st.text_area("Agentes Nocivos / Riscos Alegados", value=v_age, height=calcula_altura(v_age, 100))
                    v_ped = p_atual.get("pedidos_tecnicos", "")
                    p_atual["pedidos_tecnicos"] = st.text_area("Pedidos Técnicos (Grau, Enquadramento, PPP)", value=v_ped, height=calcula_altura(v_ped, 100))
                    v_pre = p_atual.get("preliminares_periciais", "")
                    p_atual["preliminares_periciais"] = st.text_area("Preliminares Periciais (Contestação)", value=v_pre, height=calcula_altura(v_pre, 100))
                    v_def = p_atual.get("defesa_merito_sst", "")
                    p_atual["defesa_merito_sst"] = st.text_area("Defesa de Mérito SST (Contestação)", value=v_def, height=calcula_altura(v_def, 120))

                    if st.form_submit_button("💾 Salvar Contrato e Sínteses"):
                        salvar_processo_com_feedback(
                            processo_id_selecionado,
                            p_atual,
                            "✅ Contrato e sínteses salvos!",
                        )

            with tab3:
                st.markdown("### 3. Vistoria & Análise de Documentos")
                with st.form(f"ft_3_{processo_id_selecionado}"):
                    col5, col6, col7 = st.columns(3)
                    with col5: p_atual["fase_processual"] = st.text_input("Fase Processual Atual", value=p_atual.get("fase_processual", ""))
                    with col6: p_atual["campo_data"] = st.text_input("Data da Vistoria", value=p_atual.get("campo_data", ""))
                    with col7: p_atual["campo_horario"] = st.text_input("Horário da Vistoria", value=p_atual.get("campo_horario", ""))
                    
                    p_atual["local_diligencia"] = st.text_input("Local / Endereço da Diligência", value=p_atual.get("local_diligencia", ""))
                    
                    v_ltc = p_atual.get("doc_ltcat", "")
                    p_atual["doc_ltcat"] = st.text_area("Análise do LTCAT", value=v_ltc, height=calcula_altura(v_ltc, 90))
                    v_lau = p_atual.get("doc_laudo", "")
                    p_atual["doc_laudo"] = st.text_area("Análise de Laudos Prévios / Paradigmas", value=v_lau, height=calcula_altura(v_lau, 90))
                    v_ppp = p_atual.get("doc_ppp", "")
                    p_atual["doc_ppp"] = st.text_area("Análise do PPP (Agentes, Responsáveis, EPI)", value=v_ppp, height=calcula_altura(v_ppp, 90))
                    v_pgr = p_atual.get("doc_pgr", "")
                    p_atual["doc_pgr"] = st.text_area("Análise do PGR / PPRA / PCMAT", value=v_pgr, height=calcula_altura(v_pgr, 90))
                    v_dos = p_atual.get("doc_os", "")
                    p_atual["doc_os"] = st.text_area("Ordens de Serviço e Treinamentos", value=v_dos, height=calcula_altura(v_dos, 90))
                    v_aso = p_atual.get("doc_asos", "")
                    p_atual["doc_asos"] = st.text_area("ASOs / PCMSO (Aptidão e Riscos)", value=v_aso, height=calcula_altura(v_aso, 90))
                    v_out = p_atual.get("doc_outros", "")
                    p_atual["doc_outros"] = st.text_area("Outros Documentos Relevantes (FISPQs, etc.)", value=v_out, height=calcula_altura(v_out, 90))

                    if st.form_submit_button("💾 Salvar Análise de Documentos"):
                        salvar_processo_com_feedback(
                            processo_id_selecionado,
                            p_atual,
                            "✅ Análise de documentos salva!",
                        )

            with tab4:
                st.markdown("### 4. Quadro de Fornecimento de EPIs")
                epis_atuais = p_atual.get("quadro_epis", [])
                df_epis = pd.DataFrame(epis_atuais, columns=["descricao", "ca", "data_entrega", "obs"])
                edited_df = st.data_editor(
                    df_epis,
                    column_config={
                        "descricao": st.column_config.TextColumn("Descrição do EPI fornecido", width="large", required=True),
                        "ca": st.column_config.TextColumn("C.A.", width="small"),
                        "data_entrega": st.column_config.TextColumn("Data Entrega", width="medium"),
                        "obs": st.column_config.TextColumn("Observações", width="large"),
                    },
                    num_rows="dynamic", width="stretch", height=350,
                    key=f"epis_{processo_id_selecionado}",
                )

                if st.button("💾 Salvar Tabela EPIs", key=f"btn_st_{processo_id_selecionado}"):
                    df_clean = edited_df.fillna("")
                    df_clean = df_clean[df_clean["descricao"].astype(str).str.strip() != ""]
                    p_atual["quadro_epis"] = df_clean.to_dict('records')
                    salvar_processo_com_feedback(
                        processo_id_selecionado,
                        p_atual,
                        "✅ EPIs salvos!",
                    )

                with st.form(f"ft_4_epi_{processo_id_selecionado}"):
                    v_epi2 = p_atual.get("analise_epis_critica", "")
                    p_atual["analise_epis_critica"] = st.text_area("Síntese e Análise Crítica de EPIs", value=v_epi2, height=calcula_altura(v_epi2, 180))
                    if st.form_submit_button("💾 Salvar Síntese Crítica"):
                        salvar_processo_com_feedback(
                            processo_id_selecionado,
                            p_atual,
                            "✅ Síntese crítica salva!",
                        )

            with tab5:
                st.markdown("### 5. Quesitos Formulados para a Perícia")
                with st.form(f"ft_5_{processo_id_selecionado}"):
                    v_q1 = p_atual.get("quesitos_juizo", "")
                    p_atual["quesitos_juizo"] = st.text_area("9.1. Quesitos do Juízo", value=v_q1, height=calcula_altura(v_q1, 150))
                    v_q2 = p_atual.get("quesitos_autor", "")
                    p_atual["quesitos_autor"] = st.text_area("9.2. Quesitos do Reclamante (Autor/Autora)", value=v_q2, height=calcula_altura(v_q2, 250))
                    v_q3 = p_atual.get("quesitos_reu", "")
                    p_atual["quesitos_reu"] = st.text_area("9.3. Quesitos da Reclamada (Ré / Empresa)", value=v_q3, height=calcula_altura(v_q3, 250))

                    if st.form_submit_button("💾 Salvar Quesitos Literais"):
                        salvar_processo_com_feedback(
                            processo_id_selecionado,
                            p_atual,
                            "✅ Quesitos salvos!",
                        )

elif opcao == "🚜 Diligência de Campo & Fotos":
    if not db_processos or processo_id_selecionado == "Nenhum caso cadastrado":
        st.warning("Cadastre ou selecione um caso no menu lateral.")
    else:
        caso_id = processo_id_selecionado
        p_salvo = copy.deepcopy(db_processos[caso_id])
        p_atual = db_processos[caso_id]
        rascunho_campo = st.session_state.campo_rascunhos.get(caso_id, {})
        render_caso_ativo(caso_id, p_atual, "")

        def _valor_inicial_campo(chave):
            return rascunho_campo.get(chave, p_atual.get(chave, ""))

        st.markdown("### 🚜 Vistoria Pericial de Campo")
        st.caption("Campos marcados com * são obrigatórios. Você pode digitar, usar voz ou combinar os dois no mesmo campo.")
        render_field_draft_support(caso_id)

        validation_errors_key = f"campo_validation_errors_{caso_id}"
        campo_errors = st.session_state.get(validation_errors_key, {})
        remembered_fields = st.session_state.campo_valores_memorizados
        if campo_errors:
            valores_atuais = {
                field["key"]: st.session_state.get(
                    _field_widget_key(field, caso_id),
                    _valor_inicial_campo(field["key"]),
                )
                for field in FIELD_RECORD_FIELDS
            }
            campo_errors = {
                key: value
                for key, value in campo_errors.items()
                if key in validar_registro_campo(valores_atuais)
            }
            st.session_state[validation_errors_key] = campo_errors

        st.markdown("<div class='campo-shortcuts-title'>Atalhos rápidos</div>", unsafe_allow_html=True)
        shortcut_cols = st.columns(4, gap="small")
        with shortcut_cols[0]:
            if st.button("📅 Usar hoje", key=f"campo_hoje_{caso_id}", width="stretch"):
                _set_field_widget_value(
                    FIELD_RECORD_FIELDS[0],
                    caso_id,
                    datetime.now().strftime("%d/%m/%Y"),
                )
        with shortcut_cols[1]:
            if st.button("🕒 Usar agora", key=f"campo_agora_{caso_id}", width="stretch"):
                _set_field_widget_value(
                    FIELD_RECORD_FIELDS[1],
                    caso_id,
                    datetime.now().strftime("%H:%M"),
                )
        with shortcut_cols[2]:
            if st.button(
                "↺ Último local",
                key=f"campo_local_{caso_id}",
                width="stretch",
                disabled=not remembered_fields.get("local_diligencia"),
            ):
                _set_field_widget_value(
                    FIELD_RECORD_FIELDS[2],
                    caso_id,
                    remembered_fields.get("local_diligencia", ""),
                )
        with shortcut_cols[3]:
            if st.button(
                "↺ Últimos presentes",
                key=f"campo_presentes_{caso_id}",
                width="stretch",
                disabled=not remembered_fields.get("presentes_pericia"),
            ):
                _set_field_widget_value(
                    FIELD_RECORD_FIELDS[3],
                    caso_id,
                    remembered_fields.get("presentes_pericia", ""),
                )

        def _render_campo_texto(field):
            placeholder = _field_placeholder(field, caso_id)
            current_value = _valor_inicial_campo(field["key"])
            if field["widget"] == "text_input":
                p_atual[field["key"]] = st.text_input(
                    field["label"],
                    value=current_value,
                    key=_field_widget_key(field, caso_id),
                    placeholder=placeholder,
                )
            else:
                p_atual[field["key"]] = st.text_area(
                    field["label"],
                    value=current_value,
                    height=calcula_altura(current_value, field["min_height"]),
                    key=_field_widget_key(field, caso_id),
                    placeholder=placeholder,
                )
            render_field_error(campo_errors.get(field["key"]))
            render_voice_dictation_control(
                placeholder,
                field["voice_label"],
                _field_widget_key(field, caso_id),
            )

        date_field, time_field = FIELD_RECORD_FIELDS[0], FIELD_RECORD_FIELDS[1]
        col_c1, col_c2 = st.columns(2, gap="small")
        with col_c1:
            _render_campo_texto(date_field)
        with col_c2:
            _render_campo_texto(time_field)

        for field in FIELD_RECORD_FIELDS[2:]:
            _render_campo_texto(field)

        if campo_errors:
            st.warning("Revise os campos destacados antes de salvar os textos da vistoria.")

        if st.button("💾 Salvar Textos da Vistoria", type="primary", width="stretch",
                     key=f"campo_salvar_{caso_id}"):
            erros = validar_registro_campo(p_atual)
            if erros:
                st.session_state[validation_errors_key] = erros
                st.rerun()

            st.session_state[validation_errors_key] = {}
            candidato_textos = copy.deepcopy(p_salvo)
            candidato_textos.update({key: p_atual.get(key, "") for key in FIELD_RECORD_KEYS})
            if salvar_processo(caso_id, candidato_textos):
                p_salvo = candidato_textos
                p_atual["_revision"] = candidato_textos.get("_revision")
                st.session_state.campo_valores_memorizados = atualizar_campos_memorizados(
                    st.session_state.campo_valores_memorizados,
                    p_atual,
                )
                st.toast(f"✅ Textos da vistoria guardados na Nuvem (caso {caso_id})!", icon="💾")
            else:
                st.error(f"❌ Os textos da vistoria NÃO foram salvos no caso {caso_id}; continuam preservados nesta sessão.")

        campos_alterados = {
            key: p_atual.get(key, "")
            for key in FIELD_RECORD_KEYS
            if str(p_atual.get(key, "") or "") != str(p_salvo.get(key, "") or "")
        }
        if campos_alterados:
            st.session_state.campo_rascunhos[caso_id] = campos_alterados
            st.warning(
                f"Os textos da vistoria do caso {caso_id} têm alterações NÃO salvas na nuvem. "
                "Elas ficam preservadas nesta sessão apenas para este caso; use Salvar ou descarte."
            )
            st.button(
                "↩️ Descartar alterações não salvas dos textos",
                key=f"campo_descartar_{caso_id}",
                on_click=_descartar_rascunho_campo,
                args=(caso_id,),
            )
        else:
            st.session_state.campo_rascunhos.pop(caso_id, None)

        st.markdown("<br>---<br>", unsafe_allow_html=True)
        st.markdown("#### 📍 Captura Rápida de GPS")
        gps_caso = st.text_input(
            f"Coordenada GPS Atual (caso {caso_id}):",
            value=st.session_state.gps_por_caso.get(caso_id, ""),
            key=f"gps_campo_{caso_id}",
            placeholder="GPS_TARGET_FIELD",
        )
        st.session_state.gps_por_caso[caso_id] = gps_caso

        components.html("""
            <div style="background: #E2E8F0; padding: 12px; border-radius: 8px; text-align: center; font-family: sans-serif;">
                <button onclick="getGPS()" style="background-color: #1B365D; color: white; border: none; padding: 12px 20px; border-radius: 6px; font-weight: bold; font-size: 15px; cursor: pointer; box-shadow: 0 4px 6px rgba(0,0,0,0.1);">📍 Obter GPS do Tablet Agora</button>
                <p id="gps_status" style="margin-top: 8px; font-weight: bold; color: #1B365D; font-size: 13px;"></p>
                <script>
                function injectGPS(coordStr) {
                    const mainInputs = window.parent.document.querySelectorAll('input[placeholder="GPS_TARGET_FIELD"]');
                    mainInputs.forEach(input => {
                        let nativeInputValueSetter = Object.getOwnPropertyDescriptor(window.parent.HTMLInputElement.prototype, "value").set;
                        nativeInputValueSetter.call(input, coordStr);
                        input.dispatchEvent(new Event('input', { bubbles: true })); input.dispatchEvent(new Event('change', { bubbles: true }));
                        input.dispatchEvent(new KeyboardEvent('keydown', {key: 'Enter', code: 'Enter', keyCode: 13, which: 13, bubbles: true}));
                    });
                    const photoInputs = window.parent.document.querySelectorAll('input[placeholder="GPS_FOTO"]');
                    photoInputs.forEach(input => {
                        if(!input.value || input.value === "") { 
                            let nativeInputValueSetter = Object.getOwnPropertyDescriptor(window.parent.HTMLInputElement.prototype, "value").set;
                            nativeInputValueSetter.call(input, coordStr);
                            input.dispatchEvent(new Event('input', { bubbles: true })); input.dispatchEvent(new Event('change', { bubbles: true }));
                        }
                    });
                }
                function getGPS() {
                    var st = document.getElementById("gps_status"); st.innerHTML = "Buscando satélites...";
                    if (navigator.geolocation) {
                        navigator.geolocation.getCurrentPosition(
                            function(position) {
                                var coordStr = position.coords.latitude.toFixed(6) + ", " + position.coords.longitude.toFixed(6);
                                st.innerHTML = "✅ GPS Obtido!"; injectGPS(coordStr);
                            },
                            function(error) {
                                st.innerHTML = "Tentando via rede 4G...";
                                navigator.geolocation.getCurrentPosition(
                                    function(position) {
                                        var coordStr = position.coords.latitude.toFixed(6) + ", " + position.coords.longitude.toFixed(6);
                                        st.innerHTML = "✅ GPS Obtido!"; injectGPS(coordStr);
                                    },
                                    function(err2) { st.innerHTML = "❌ Erro de GPS."; },
                                    {timeout: 10000, enableHighAccuracy: false, maximumAge: 60000}
                                );
                            }, {timeout: 8000, enableHighAccuracy: true, maximumAge: 0}
                        );
                    } else { st.innerHTML = "Geolocalização não suportada."; }
                }
                setInterval(function() {
                    const textareas = window.parent.document.querySelectorAll("textarea");
                    textareas.forEach(ta => {
                        if(ta.scrollHeight > ta.clientHeight) { ta.style.height = 'auto'; ta.style.height = ta.scrollHeight + 'px'; }
                        ta.addEventListener('input', function() { this.style.height = 'auto'; this.style.height = this.scrollHeight + 'px'; });
                    });
                }, 1000);
                </script>
            </div>
        """, height=95)

        st.markdown("<br>", unsafe_allow_html=True)
        st.markdown("#### 📸 Fotos do Caso")

        fotos_salvas = list(p_salvo.get("campo_fotos", []) or [])
        pendentes = list(st.session_state.fotos_pendentes.get(caso_id, []))
        versao_key = f"fotos_widget_versao_{caso_id}"
        versao_fotos = st.session_state.get(versao_key, 0)
        aviso_fotos_key = f"fotos_aviso_{caso_id}"
        painel_fotos_key = _painel_fotos_key(caso_id)
        if painel_fotos_key not in st.session_state:
            # Caso sem fotos já abre a área de adição; com fotos, ela é aberta pelo botão.
            st.session_state[painel_fotos_key] = not fotos_salvas
        if st.session_state.get(aviso_fotos_key):
            st.info(st.session_state.pop(aviso_fotos_key))

        resumo_fotos = resumo_armazenamento(p_salvo, pendentes)
        with st.container(border=True):
            st.markdown(
                f"**📍 Onde ficam as fotos:** fotos novas do caso **{caso_id}** ficam no **Firebase Storage privado**. "
                f"No Firestore › coleção `processos` › documento `{caso_id}` › campo `campo_fotos` "
                "ficam apenas referências e metadados das fotos novas. Fotos legadas continuam embutidas ou no caminho original. "
                "Não existe pasta nova no aparelho: depois de salvas, as fotos acompanham o caso "
                "em qualquer dispositivo e entram no Documento Word final."
            )
            col_m1, col_m2, col_m3 = st.columns(3)
            col_m1.metric("☁️ Salvas na nuvem", resumo_fotos["salvas"])
            col_m2.metric("⏳ Pendentes (só nesta sessão)", resumo_fotos["pendentes"])
            col_m3.metric(
                "💾 Documento Firestore (estimativa)",
                f"{resumo_fotos['bytes_usados'] / 1024:.0f} / {resumo_fotos['bytes_limite'] / 1024:.0f} KB",
            )
            st.progress(
                resumo_fotos["fracao_usada"],
                text=f"{resumo_fotos['fracao_usada'] * 100:.0f}% do orçamento do documento (~1 MiB; sem os bytes das fotos novas)",
            )
            st.caption(
                f"Storage: {resumo_fotos['storage']} foto(s) · Legadas: {resumo_fotos['legadas']} foto(s). "
                "O documento ainda contém textos, metadados e base64 legado. "
                "Storage pode ter custos de armazenamento, operações e downloads. Use lotes pequenos; "
                "cada foto aceita até 10 MiB após compressão. "
                "Fotos pendentes ficam apenas na memória desta sessão do navegador até serem gravadas "
                "no caso; se a página for fechada ou recarregada antes disso, elas se perdem."
            )
        if resumo_fotos["fracao_usada"] > 0.85:
            st.warning(
                "O caso está próximo do limite de armazenamento por documento do banco. "
                "Mesmo com fotos novas no Storage, referências, textos e fotos legadas precisam caber. "
                "Revise o conteúdo do caso; fotos antigas não são migradas automaticamente."
            )

        st.markdown("---")
        st.markdown(f"#### 🖼️ Gerenciar Fotos ({len(fotos_salvas)} fotos salvas na nuvem):")

        # Botão grande, fora dos cartões das imagens, que abre/fecha a área de adição.
        painel_fotos_aberto = st.session_state[painel_fotos_key]
        st.markdown(
            f"<style>.st-key-fotos_adicionar_{caso_id} button {{min-height: 3.5rem;}}"
            f".st-key-fotos_adicionar_{caso_id} button p {{font-size: 1.2rem; font-weight: 700;}}</style>",
            unsafe_allow_html=True,
        )
        st.button(
            "✖️ Fechar área de adição de fotos" if painel_fotos_aberto else "➕ Adicionar novas fotos",
            key=f"fotos_adicionar_{caso_id}",
            type="secondary" if painel_fotos_aberto else "primary",
            width="stretch",
            on_click=_alternar_painel_fotos, args=(caso_id,),
            help=f"Envia fotos ao Storage e vincula cada lote ao caso {caso_id} no Firestore.",
        )

        img_camera = None
        fotos_upload = None
        if painel_fotos_aberto:
            with st.container(border=True):
                st.markdown(f"##### ➕ Adicionar fotos ao caso {caso_id}")
                st.caption(
                    "Selecione uma ou várias fotos de uma vez. Cada lote é gravado automaticamente "
                    f"no Storage e vinculado ao caso {caso_id}; após salvar, esta área continua aberta "
                    "para o próximo lote — as fotos já salvas permanecem na galeria abaixo."
                )
                # Widgets de captura com chave por caso + versão: após mover os arquivos para a
                # fila de pendentes a versão muda, liberando a captura para o próximo ciclo.
                fotos_upload = st.file_uploader(
                    "📸 CÂMERA DO TABLET (Traseira c/ Zoom) ou Galeria", type=["jpg", "jpeg", "png"],
                    accept_multiple_files=True, key=f"fotos_upload_{caso_id}_{versao_fotos}",
                )
                img_camera = st.camera_input(
                    "📷 Câmera Web (Apenas para Computador)", key=f"camera_{caso_id}_{versao_fotos}"
                )

        arquivos_novos = ([img_camera] if img_camera is not None else []) + list(fotos_upload or [])
        salvar_pendentes_agora = False
        if arquivos_novos:
            novas_fotos = [
                criar_foto(comprimir_imagem(arquivo.getvalue()), gps_caso)
                for arquivo in arquivos_novos
            ]
            pendentes, adicionadas, duplicadas = adicionar_pendentes(fotos_salvas, pendentes, novas_fotos)
            st.session_state.fotos_pendentes[caso_id] = pendentes
            st.session_state[versao_key] = versao_fotos + 1
            st.session_state[painel_fotos_key] = True
            if duplicadas:
                st.session_state[aviso_fotos_key] = (
                    f"{duplicadas} foto(s) ignorada(s) por já estarem cadastradas neste caso."
                )
            if adicionadas:
                salvar_pendentes_agora = True
            else:
                st.rerun()

        if pendentes:
            st.warning(
                f"⚠️ {len(pendentes)} foto(s) aguardando gravação no caso {caso_id} — "
                "ainda NÃO confirmadas no caso. Um upload pode já ter ocorrido; a fila só sai após "
                "confirmar o Firestore. Elas não são enviadas para outros casos."
            )
            cols_pend = st.columns(min(len(pendentes), 4))
            for indice_pend, foto_pend in enumerate(pendentes):
                with cols_pend[indice_pend % len(cols_pend)]:
                    st.image(resolver_bytes_foto(foto_pend), width=120)
            col_p1, col_p2, col_p3 = st.columns(3)
            with col_p1:
                if st.button("🔁 Tentar salvar fotos pendentes", key=f"fotos_reenviar_{caso_id}"):
                    salvar_pendentes_agora = True
            with col_p2:
                st.button(
                    "🔄 Recarregar versão da nuvem", key=f"fotos_recarregar_{caso_id}",
                    on_click=_recarregar_caso, args=(caso_id,),
                    help="Use se o caso foi alterado em outro dispositivo. As fotos pendentes são mantidas.",
                )
            with col_p3:
                st.button(
                    "🗑️ Descartar fotos pendentes", key=f"fotos_descartar_{caso_id}",
                    on_click=_descartar_fotos_pendentes, args=(caso_id,),
                )

        if salvar_pendentes_agora:
            # Recarregar a nuvem pode revelar fotos já salvas por outra tentativa.
            pendentes, _, _ = adicionar_pendentes(fotos_salvas, [], pendentes)
            st.session_state.fotos_pendentes[caso_id] = pendentes
            enviadas = []
            erro_fotos = ""
            if caso_id != st.session_state.get("processo_ativo"):
                erro_fotos = "Gravação bloqueada: o destino não é o caso ativo."
            else:
                for foto in pendentes:
                    try:
                        referencia = metadados_storage(caso_id, foto)
                        cabem, _ = separar_por_limite(
                            {**p_salvo, "campo_fotos": fotos_salvas + enviadas}, [referencia],
                        )
                        if not cabem:
                            erro_fotos = (
                                f"Limite de armazenamento do caso {caso_id} atingido: "
                                "os metadados não cabem no documento Firestore (~1 MiB). "
                                "As fotos restantes continuam pendentes; revise textos e fotos legadas."
                            )
                            break
                        enviadas.append(enviar_foto(caso_id, foto))
                    except Exception as exc:
                        logger.exception("Falha no upload de fotos do caso %s.", caso_id)
                        erro_fotos = f"Falha no Storage: {exc}. As fotos restantes continuam pendentes."
                        break
            if enviadas:
                candidato_fotos = copy.deepcopy(p_salvo)
                candidato_fotos["campo_fotos"] = fotos_salvas + enviadas
                if salvar_processo(caso_id, candidato_fotos):
                    restantes = pendentes[len(enviadas):]
                    if restantes:
                        st.session_state.fotos_pendentes[caso_id] = restantes
                        st.session_state[aviso_fotos_key] = erro_fotos
                    else:
                        st.session_state.fotos_pendentes.pop(caso_id, None)
                    st.toast(
                        f"✅ {len(enviadas)} foto(s) no Storage e vinculada(s) ao caso {caso_id}!",
                        icon="📸",
                    )
                    st.rerun()
                else:
                    st.error(
                        "❌ As fotos NÃO foram vinculadas ao caso. Os uploads já feitos serão reutilizados. "
                        "Elas continuam pendentes nesta sessão; use “Tentar salvar fotos pendentes”."
                    )
            elif erro_fotos:
                st.error(erro_fotos)
            else:
                st.session_state.fotos_pendentes.pop(caso_id, None)
                st.rerun()

        col_gps, col_limpar = st.columns(2)
        with col_gps:
            if st.button("⚡ Salvar GPS em Todas as Fotos", key=f"fotos_gps_todas_{caso_id}"):
                if gps_caso and fotos_salvas:
                    candidato_fotos = copy.deepcopy(p_salvo)
                    for f_dict in candidato_fotos["campo_fotos"]: f_dict["gps"] = gps_caso
                    if salvar_processo(caso_id, candidato_fotos):
                        st.rerun()
                    else:
                        st.error("❌ Não foi possível atualizar o GPS das fotos.")
        with col_limpar:
            confirmar_limpeza = st.checkbox(
                f"Confirmo excluir todas as fotos do caso {caso_id}",
                key=f"fotos_limpar_confirmar_{caso_id}",
                disabled=not fotos_salvas,
            )
            if st.button("🗑️ Limpar Todas Fotos", key=f"fotos_limpar_{caso_id}",
                         disabled=not (confirmar_limpeza and fotos_salvas)):
                candidato_fotos = copy.deepcopy(p_salvo)
                candidato_fotos["campo_fotos"] = []
                if salvar_processo(caso_id, candidato_fotos):
                    st.rerun()
                else:
                    st.error("❌ Não foi possível limpar as fotos.")
        if not fotos_salvas:
            st.info("Nenhuma foto cadastrada.")
        else:
            for idx, (foto_id, foto_dict) in enumerate(zip(ids_fotos(fotos_salvas), fotos_salvas)):
                with st.container():
                    col_img, col_dados = st.columns([1, 2])
                    with col_img:
                        try:
                            st.image(resolver_bytes_foto(foto_dict), width=220)
                        except Exception:
                            logger.exception("Falha ao exibir foto do caso %s.", caso_id)
                            st.warning("Imagem indisponível. Verifique a conexão/configuração do Storage ou a foto legada.")
                    with col_dados:
                        nova_legenda = st.text_input(
                            f"Legenda {idx+1}",
                            value=foto_dict.get("legenda", ""),
                            key=f"lg_{caso_id}_{foto_id}",
                        )
                        novo_gps = st.text_input(
                            f"GPS {idx+1}",
                            value=foto_dict.get("gps", ""),
                            key=f"cg_{caso_id}_{foto_id}",
                            placeholder="GPS_FOTO",
                        )

                        c_salvar, c_del = st.columns(2)
                        with c_salvar:
                            if st.button(f"💾 Atualizar {idx+1}", key=f"foto_atualizar_{caso_id}_{foto_id}"):
                                candidato_fotos = copy.deepcopy(p_salvo)
                                candidato_fotos["campo_fotos"][idx]["legenda"] = nova_legenda
                                candidato_fotos["campo_fotos"][idx]["gps"] = novo_gps
                                if salvar_processo(caso_id, candidato_fotos):
                                    st.rerun()
                                else:
                                    st.error("❌ Não foi possível atualizar os dados da foto.")
                        with c_del:
                            if st.button(f"🗑️ Excluir {idx+1}", key=f"foto_excluir_{caso_id}_{foto_id}"):
                                candidato_fotos = copy.deepcopy(p_salvo)
                                candidato_fotos["campo_fotos"].pop(idx)
                                if salvar_processo(caso_id, candidato_fotos):
                                    st.rerun()
                                else:
                                    st.error("❌ Não foi possível excluir a foto.")
                    st.markdown("---")

elif opcao == "🗑️ Excluir Processo":
    st.markdown("### Excluir Caso / Processo")
    if not db_processos or processo_id_selecionado == "Nenhum caso cadastrado":
        st.warning("Nenhum caso disponível.")
    else:
        p_excluir = db_processos[processo_id_selecionado]
        st.error(f"⚠️ Atenção: Você está prestes a excluir o caso **{processo_id_selecionado}**.")
        
        if st.session_state.confirmar_exclusao_target != processo_id_selecionado:
            if st.button("🗑️ Solicitar Exclusão Definitiva"):
                st.session_state.confirmar_exclusao_target = processo_id_selecionado
                st.session_state.confirmar_exclusao_input_nonce += 1
                st.rerun()
        else:
            st.warning("🚨 TEM CERTEZA ABSOLUTA?")
            st.caption(
                f"Confirme digitando o ID exato do caso: `{processo_id_selecionado}`"
            )
            confirmacao_id = st.text_input(
                "Digite o ID para confirmar a exclusão permanente:",
                key=f"confirmar_exclusao_input_{st.session_state.confirmar_exclusao_input_nonce}",
            ).strip()
            col_sim, col_nao = st.columns(2)
            with col_sim:
                if st.button("🔴 SIM, EXCLUIR PERMANENTEMENTE"):
                    if confirmacao_id != processo_id_selecionado:
                        st.error("❌ Confirmação inválida. A exclusão foi cancelada.")
                    else:
                        if excluir_processo(processo_id_selecionado):
                            st.session_state.confirmar_exclusao_target = ""
                            st.session_state.confirmar_exclusao_input_nonce += 1
                            st.session_state.campo_rascunhos.pop(processo_id_selecionado, None)
                            st.session_state.fotos_pendentes.pop(processo_id_selecionado, None)
                            st.session_state.gps_por_caso.pop(processo_id_selecionado, None)
                            st.session_state.processo_ativo = None
                            st.session_state.menu_opcao = "➕ Novo Processo / Caso"
                            st.toast(f"Caso {processo_id_selecionado} excluído. Selecione o próximo caso no menu lateral.")
                            st.rerun()
            with col_nao:
                if st.button("❌ Cancelar"):
                    st.session_state.confirmar_exclusao_target = ""
                    st.session_state.confirmar_exclusao_input_nonce += 1
                    st.rerun()

elif opcao == "📄 Gerar Documento Word Final":
    st.markdown("### Geração do Documento Final em Word (.docx)")
    if not db_processos or processo_id_selecionado == "Nenhum caso cadastrado":
        st.warning("Selecione um caso válido no menu lateral.")
    else:
        p = db_processos[processo_id_selecionado]
        is_prev = "Previdenciário" in p.get("modulo_atuacao", "")
        st.markdown(
            f"**Caso:** {html.escape(str(processo_id_selecionado))} — "
            f"**Segurado:** {html.escape(str(p.get('reclamante_nome', '')))}"
        )
        
        st.markdown("<br>", unsafe_allow_html=True)
        erros_revisao = final_errors(p) if is_prev else []
        p, erros_documento = _normalizar_dados_documento(p)
        if is_prev:
            erros_documento.extend(erros_revisao)
            st.caption("LTCAT e dados de apoio ao PPP: geração manual; não é PPP oficial assinado.")
            if st.button("📋 Gerar relatório pré-pericial / rascunho",
                         key=f"extra_{processo_id_selecionado}_preliminar"):
                try:
                    draft, warnings = build_extrajudicial_bytes(p, preliminary=True)
                    for warning in warnings:
                        st.warning(warning)
                    st.download_button("📥 Baixar relatório pré-pericial", draft,
                                       file_name=f"Rascunho_{processo_id_selecionado}.docx",
                                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                except (ValueError, TypeError) as exc:
                    st.error(f"Não foi possível gerar o rascunho: {exc}")
        for erro_documento in erros_documento:
            st.error(erro_documento)
        if st.button(
            "📥 Gerar LTCAT revisado + apoio PPP (.docx)" if is_prev else "📥 Gerar e Baixar Documento Oficial (.docx)",
            disabled=bool(erros_documento),
            key=f"gerar_documento_{processo_id_selecionado}",
        ):
            doc = Document()
            for section in doc.sections:
                section.top_margin = Cm(3.0)
                section.left_margin = Cm(3.0)
                section.bottom_margin = Cm(2.0)
                section.right_margin = Cm(2.0)

            def adicionar_linha_vazia():
                p_v = doc.add_paragraph()
                p_v.paragraph_format.line_spacing = 1.5
                return p_v

            def add_bullet_inline(doc_obj, title, text):
                p_b = doc_obj.add_paragraph()
                p_b.paragraph_format.line_spacing = 1.5
                p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                r_title = p_b.add_run(f"• {title}: ")
                r_title.bold = True
                r_title.font.name = 'Abadi'
                r_title.font.size = Pt(11)
                r_text = p_b.add_run(f"{text if text else '[Informação não localizada]'}")
                r_text.font.name = 'Abadi'
                r_text.font.size = Pt(11)

            def add_topic_block(doc_obj, title, text):
                p_title = doc_obj.add_paragraph()
                p_title.paragraph_format.line_spacing = 1.5
                p_title.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                r_title = p_title.add_run(f"• {title}:")
                r_title.bold = True
                r_title.font.name = 'Abadi'
                r_title.font.size = Pt(11)
                if not text: text = "[Informação não localizada]"
                for linha in text.split("\n"):
                    if linha.strip():
                        p_text = doc_obj.add_paragraph()
                        p_text.paragraph_format.line_spacing = 1.5
                        p_text.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        r_text = p_text.add_run(linha.strip())
                        r_text.font.name = 'Abadi'
                        r_text.font.size = Pt(11)

            def adicionar_titulo(texto, level=2):
                h = doc.add_heading(level=level)
                h.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                h.paragraph_format.space_before = Pt(18)
                h.paragraph_format.space_after = Pt(6)
                h.paragraph_format.line_spacing = 1.5
                run = h.add_run(texto)
                run.font.name = 'Abadi'
                run.font.size = Pt(14) if level <= 2 else Pt(12)
                run.bold = True
                run.font.color.rgb = RGBColor(27, 54, 93)

            def definir_largura_tabela_100(tabela):
                tblPr = tabela._tbl.tblPr
                tblW = OxmlElement('w:tblW')
                tblW.set(qn('w:w'), '5000')
                tblW.set(qn('w:type'), 'pct')
                tblPr.append(tblW)

            logo_doc = None
            if os.path.exists(LOGO_FILE): logo_doc = LOGO_FILE
            if logo_doc:
                p_logo = doc.add_paragraph()
                p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_logo.add_run().add_picture(logo_doc, width=Inches(2.0))
                adicionar_linha_vazia()

            t_main = doc.add_paragraph()
            t_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
            t_main.paragraph_format.line_spacing = 1.5
            
            titulo_doc = "LAUDO TÉCNICO DE CONDIÇÕES AMBIENTAIS DO TRABALHO (LTCAT)" if is_prev else "PRÉ-RELATÓRIO DE ANÁLISE PROCESSUAL E PERICIAL"
            r_t = t_main.add_run(titulo_doc)
            r_t.bold = True
            r_t.font.name = 'Abadi'
            r_t.font.size = Pt(15)
            r_t.font.color.rgb = RGBColor(27, 54, 93)
            
            p_papel = doc.add_paragraph()
            p_papel.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_papel.paragraph_format.line_spacing = 1.5
            r_p1 = p_papel.add_run(f"Módulo: {p.get('modulo_atuacao', '')}\n")
            r_p2 = p_papel.add_run(f"Profissional: {p.get('papel_profissional', 'Perito / Consultor')}")
            for r in [r_p1, r_p2]:
                r.bold = True
                r.font.name = 'Abadi'
                r.font.size = Pt(11)

            if is_prev:
                append_extrajudicial(doc, p)

            else:
                adicionar_titulo("1. IDENTIFICAÇÃO DO PROCESSO", level=2)
                dados_ident = [
                    ("Número do Processo", p.get('processo_num', '')), ("Órgão Julgador / Vara", p.get('orgao_julgador', '')),
                    ("Data de Autuação (Ajuizamento)", p.get('data_autuacao', '')), ("Valor da Causa", p.get('valor_causa', '')),
                    ("Rito Processual", p.get('rito_processual', ''))
                ]
                t_id = doc.add_table(rows=len(dados_ident) + 1, cols=2)
                definir_largura_tabela_100(t_id)
                t_id.rows[0].cells[0].text = "Parâmetro"
                t_id.rows[0].cells[1].text = "Descrição"
                for idx, (k, v) in enumerate(dados_ident, start=1):
                    t_id.rows[idx].cells[0].text = k
                    t_id.rows[idx].cells[1].text = v or "[Informação não localizada]"

                for i, row in enumerate(t_id.rows):
                    for cell in row.cells:
                        tcPr = cell._tc.get_or_add_tcPr()
                        for p_cell in cell.paragraphs:
                            p_cell.paragraph_format.line_spacing = 1.15
                            for run in p_cell.runs:
                                run.font.name = 'Abadi'
                                run.font.size = Pt(9)
                        if i == 0:
                            shd = parse_xml(r'<w:shd {} w:fill="1B365D"/>'.format(nsdecls('w')))
                            tcPr.append(shd)
                            for p_cell in cell.paragraphs:
                                for run in p_cell.runs:
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                    run.bold = True
                        else:
                            if i % 2 == 0:
                                shd = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
                                tcPr.append(shd)
                adicionar_linha_vazia()

                adicionar_titulo("2. QUALIFICAÇÃO DAS PARTES", level=2)
                add_bullet_inline(doc, "Reclamante (Autor/Autora)", f"{p.get('reclamante_nome', '')} (CPF: {p.get('reclamante_cpf', '')})")
                add_bullet_inline(doc, "Advogados", p.get('reclamante_adv', ''))
                add_bullet_inline(doc, "Reclamada (Ré/Empresa)", f"{p.get('reclamada_nome', '')} (CNPJ: {p.get('reclamada_cnpj', '')})")
                add_bullet_inline(doc, "Advogados", p.get('reclamada_adv', ''))

                adicionar_titulo("3. DADOS DO CONTRATO DE TRABALHO", level=2)
                add_bullet_inline(doc, "Data de Admissão", p.get('data_admissao', ''))
                add_bullet_inline(doc, "Status do Contrato", p.get('status_contrato', ''))
                add_bullet_inline(doc, "Período Imprescrito", p.get('periodo_imprescrito', ''))
                add_bullet_inline(doc, "Cargo(s) / Função(ões)", p.get('cargos', ''))
                add_bullet_inline(doc, "Setor / Lotação / Local", p.get('setor', ''))
                add_bullet_inline(doc, "Última Remuneração", p.get('ultima_remuneracao', ''))

                adicionar_titulo("4. SÍNTESE TÉCNICA DA PETIÇÃO INICIAL", level=2)
                add_bullet_inline(doc, "Objeto da Perícia", p.get('objeto_pericia', ''))
                add_topic_block(doc, "Atividades Descritas", p.get('atividades_inicial', ''))
                add_topic_block(doc, "Agentes Nocivos / Riscos Alegados", p.get('agentes_alegados', ''))
                add_topic_block(doc, "Pedidos Técnicos", p.get('pedidos_tecnicos', ''))

                adicionar_titulo("5. SÍNTESE TÉCNICA DA CONTESTAÇÃO", level=2)
                add_topic_block(doc, "Preliminares Periciais", p.get('preliminares_periciais', ''))
                add_topic_block(doc, "Defesa de Mérito (SST)", p.get('defesa_merito_sst', ''))

                adicionar_titulo("6. STATUS E DADOS DA VISTORIA", level=2)
                add_bullet_inline(doc, "Fase Processual Atual", p.get('fase_processual', ''))
                add_bullet_inline(doc, "Data da Vistoria", p.get('campo_data', ''))
                add_bullet_inline(doc, "Horário", p.get('campo_horario', ''))
                add_bullet_inline(doc, "Local / Endereço", p.get('local_diligencia', ''))

                adicionar_titulo("7. ANÁLISE DOS DOCUMENTOS DE SST NOS AUTOS", level=2)
                add_topic_block(doc, "LTCAT", p.get('doc_ltcat', ''))
                add_topic_block(doc, "Laudo de Insalubridade / Periculosidade", p.get('doc_laudo', ''))
                add_topic_block(doc, "PPP (Perfil Profissiográfico Previdenciário)", p.get('doc_ppp', ''))
                add_topic_block(doc, "PGR / PPRA / PCMAT", p.get('doc_pgr', ''))
                add_topic_block(doc, "Ordens de Serviço (OS) / Treinamentos", p.get('doc_os', ''))
                add_topic_block(doc, "ASOs / PCMSO", p.get('doc_asos', ''))
                add_topic_block(doc, "Outros Documentos Relevantes", p.get('doc_outros', ''))

                adicionar_titulo("8. FORNECIMENTO DE EQUIPAMENTOS DE PROTEÇÃO (EPIs)", level=2)
                quadro_epis = p.get('quadro_epis', [])
                if quadro_epis:
                    t_epi = doc.add_table(rows=len(quadro_epis) + 1, cols=4)
                    definir_largura_tabela_100(t_epi)
                    headers = ["Descrição do EPI", "C.A.", "Data de Entrega", "Observações"]
                    for h_idx, h_text in enumerate(headers): t_epi.rows[0].cells[h_idx].text = h_text
                    
                    for r_idx, epi in enumerate(quadro_epis, start=1):
                        t_epi.rows[r_idx].cells[0].text = epi.get('descricao', '')
                        t_epi.rows[r_idx].cells[1].text = epi.get('ca', '')
                        t_epi.rows[r_idx].cells[2].text = epi.get('data_entrega', '')
                        t_epi.rows[r_idx].cells[3].text = epi.get('obs', '')

                    for i, row in enumerate(t_epi.rows):
                        for cell in row.cells:
                            tcPr = cell._tc.get_or_add_tcPr()
                            for p_cell in cell.paragraphs:
                                p_cell.paragraph_format.line_spacing = 1.15
                                for run in p_cell.runs:
                                    run.font.name = 'Abadi'
                                    run.font.size = Pt(8.5)
                            if i == 0:
                                shd = parse_xml(r'<w:shd {} w:fill="1B365D"/>'.format(nsdecls('w')))
                                tcPr.append(shd)
                                for p_cell in cell.paragraphs:
                                    for run in p_cell.runs:
                                        run.font.color.rgb = RGBColor(255, 255, 255)
                                        run.bold = True
                            else:
                                if i % 2 == 0:
                                    shd = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
                                    tcPr.append(shd)
                else:
                    p_sem = doc.add_paragraph("[Nenhum EPI cadastrado para este processo]")
                    p_sem.paragraph_format.line_spacing = 1.5
                    p_sem.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    p_sem.runs[0].font.name = 'Abadi'
                    p_sem.runs[0].font.size = Pt(11)

                adicionar_linha_vazia()
                add_topic_block(doc, "Síntese e Análise Crítica de EPIs", p.get('analise_epis_critica', ''))

                adicionar_titulo("9. QUESITOS FORMULADOS PARA A PERÍCIA", level=2)
                add_topic_block(doc, "9.1. Quesitos do Juízo", p.get('quesitos_juizo', ''))
                add_topic_block(doc, "9.2. Quesitos do Reclamante", p.get('quesitos_autor', ''))
                add_topic_block(doc, "9.3. Quesitos da Reclamada", p.get('quesitos_reu', ''))

                adicionar_titulo("10. LEVANTAMENTOS DE CAMPO (DILIGÊNCIA & EVIDÊNCIAS)", level=2)
                add_topic_block(doc, "Pessoas Presentes na Vistoria", p.get('presentes_pericia', ''))
                add_topic_block(doc, "Informações prestadas pelo Autor", p.get('campo_declaracoes_autor', ''))
                add_topic_block(doc, "Informações prestadas pelo Ré", p.get('campo_declaracoes_reu', ''))
                add_topic_block(doc, "Medições Realizadas", p.get('campo_medicoes', ''))

            if p.get("campo_fotos"):
                adicionar_titulo("REGISTROS FOTOGRÁFICOS DE CAMPO", level=3)
                for idx, foto_dict in enumerate(p["campo_fotos"]):
                    if foto_dict:
                        try:
                            p_img = doc.add_paragraph()
                            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            img_stream = io.BytesIO(resolver_bytes_foto(foto_dict))
                            p_img.add_run().add_picture(img_stream, width=Inches(4.5))
                            
                            texto_legenda = foto_dict.get("legenda", "Registro fotográfico.")
                            coords_gps = foto_dict.get("gps", "")
                            if coords_gps: texto_legenda += f" ({coords_gps})"
                            
                            p_leg = doc.add_paragraph(f"Figura {idx+1}: {texto_legenda}")
                            p_leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            p_leg.paragraph_format.line_spacing = 1.5
                            for run in p_leg.runs:
                                run.font.name = 'Abadi'
                                run.font.size = Pt(9)
                                run.italic = True
                            adicionar_linha_vazia()
                        except Exception as e:
                            logger.exception("Falha ao inserir foto %s no documento Word.", idx + 1)
                            st.warning(f"Não foi possível inserir a foto {idx + 1}: {e}")

            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            st.toast("✅ Documento Word Finalizado!", icon="📄")
            st.download_button(
                label="📥 Baixar LTCAT revisado + apoio PPP (.docx)" if is_prev else "📥 Baixar Documento Oficial (.docx)",
                data=buffer,
                file_name=f"Documento_Oficial_{p.get('reclamante_nome','Caso').replace(' ','_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )

pendencias_sessao = [
    f"{caso}: textos da vistoria" for caso, campos in st.session_state.campo_rascunhos.items()
    if campos and caso in db_processos
] + [
    f"{caso}: {len(fotos)} foto(s)" for caso, fotos in st.session_state.fotos_pendentes.items()
    if fotos and caso in db_processos
]
if pendencias_sessao:
    aviso_pendencias.warning(
        "Alterações NÃO salvas nesta sessão:\n\n" + "\n".join(f"- {item}" for item in pendencias_sessao)
    )
